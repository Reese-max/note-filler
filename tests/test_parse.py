from pathlib import Path

import pytest

from note_filler.parse import Document, Paragraph, parse_note

FIXTURE = Path(__file__).parent / "fixtures" / "sample.txt"


def test_parse_txt_splits_on_blank_lines():
    doc = parse_note(str(FIXTURE))
    assert isinstance(doc, Document)
    assert doc.source_path == str(FIXTURE)
    # 三段：中間的連續空行不應產生空段
    assert len(doc.paragraphs) == 3
    assert isinstance(doc.paragraphs, tuple)
    assert [p.idx for p in doc.paragraphs] == [0, 1, 2]
    assert doc.paragraphs[0].text.startswith("行政程序法第 92 條")
    assert doc.paragraphs[1].text.startswith("訴願法第 14 條")
    assert doc.paragraphs[2].text.startswith("本筆記整理自上課講義")
    # full_text = 各段以換行 join
    assert doc.full_text == "\n".join(p.text for p in doc.paragraphs)


def test_paragraph_and_document_are_frozen():
    p = Paragraph(idx=0, text="x")
    with pytest.raises(Exception):
        p.text = "y"  # frozen dataclass 不可指派
    doc = parse_note(str(FIXTURE))
    with pytest.raises(Exception):
        doc.full_text = "tampered"  # 原文 immutable


def test_parse_docx_roundtrip(tmp_path):
    # 用 python-docx 真的產生一份多段 .docx，再 parse_note 讀回，
    # 明確落實 spec「Word 輸入」需求。
    from docx import Document as DocxDocument

    paras = [
        "行政程序法第 92 條規定行政處分之定義。",
        "訴願法第 14 條：訴願應自處分達到次日起三十日內提起。",
        "第三段：交叉驗證需 >= 2 個獨立 A/B 來源。",
    ]
    src = DocxDocument()
    for t in paras:
        src.add_paragraph(t)
    src.add_paragraph("")     # 空段
    src.add_paragraph("   ")  # 只含空白 → 皆應被略過
    out = tmp_path / "note.docx"
    src.save(str(out))

    doc = parse_note(str(out))
    assert len(doc.paragraphs) == 3                     # 空段被過濾
    assert [p.text for p in doc.paragraphs] == paras     # 文字逐段正確且未改寫
    assert [p.idx for p in doc.paragraphs] == [0, 1, 2]  # idx 從 0 連續
    assert doc.full_text == "\n".join(paras)             # full_text 換行 join
