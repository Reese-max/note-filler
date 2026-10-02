"""可開啟連結驗收測試：at_least_two_openable_links 規則。

論點區塊至少需 2 條真實可開啟連結（http/https），優先採用實際引用來源，
再補延伸閱讀。若合格來源不足，輸出【待補來源】並標示原因。
禁止拼湊 URL、禁止把待補來源與真實連結混排成可通過格式檢查的假結果。
"""
from __future__ import annotations

import json
import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from note_filler.binding_report import (
    build_binding_report,
    parse_binding_report,
)
from note_filler.correction import (
    MIN_OPENABLE_LINKS,
    CorrectionDoc,
    Segment,
    assemble_correction,
    is_openable_url,
)
from note_filler.export import to_docx, to_json, to_markdown
from note_filler.gap import Gap
from note_filler.llm import FakeLLM
from note_filler.parse import Document, Paragraph, parse_note
from note_filler.pipeline import run_pipeline
from note_filler.retrieve.models import Source
from note_filler.verify import cross_validate
from note_filler.write import WrittenSupplement
from test_pipeline import FakeLaw, FakeTwinkle


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _doc(path: str = "input/note.txt") -> Document:
    return Document(
        source_path=path,
        paragraphs=(Paragraph(0, "原稿逐字保留。"),),
        full_text="原稿逐字保留。",
    )


def _source(id: str, title: str, level: str = "A", url: str | None = None) -> Source:
    return Source(
        id=id,
        title=title,
        url=url,
        level=level,
        content=f"{title} 內容",
        fetched_date="2026-07-25",
        doc_date=None,
        distance=0.5,
    )


def _note_fixture(tmp_path: Path, content: str) -> str:
    p = tmp_path / "note.docx"
    d = DocxDocument()
    for para in content.split("\n\n"):
        if para.strip():
            d.add_paragraph(para.strip())
    d.save(str(p))
    return str(p)


# ---------------------------------------------------------------------------
# is_openable_url 單元測試
# ---------------------------------------------------------------------------

class TestIsOpenableUrl:
    """is_openable_url 單元測試：真實可開啟連結判定。"""

    def test_https_url(self):
        assert is_openable_url("https://example.com") is True

    def test_http_url(self):
        assert is_openable_url("http://example.com") is True

    def test_https_with_path(self):
        assert is_openable_url("https://law.moj.gov.tw/LawClass/LawAll.aspx?pcode=D0050001") is True

    def test_ftp_url_rejected(self):
        assert is_openable_url("ftp://example.com") is False

    def test_empty_string(self):
        assert is_openable_url("") is False

    def test_whitespace_only(self):
        assert is_openable_url("   ") is False

    def test_none(self):
        assert is_openable_url(None) is False

    def test_non_string(self):
        assert is_openable_url(123) is False

    def test_no_scheme(self):
        assert is_openable_url("example.com") is False

    def test_no_netloc(self):
        assert is_openable_url("https://") is False

    def test_relative_path(self):
        assert is_openable_url("/path/to/file") is False

    def test_url_with_port(self):
        assert is_openable_url("https://example.com:8080/path") is True

    def test_url_with_query(self):
        assert is_openable_url("https://example.com?q=test&page=1") is True

    def test_url_with_fragment(self):
        assert is_openable_url("https://example.com#section") is True


# ---------------------------------------------------------------------------
# Segment 可開啟連結欄位測試
# ---------------------------------------------------------------------------

class TestSegmentOpenableLinksFields:
    """Segment 資料結構必須包含可開啟連結三欄。"""

    def test_segment_has_fields(self):
        seg = Segment(
            type="supplement",
            text="test",
            anchor_idx=None,
            sources=[],
            confidence="pending_evidence",
        )
        assert hasattr(seg, "openable_links_count")
        assert hasattr(seg, "openable_links_status")
        assert hasattr(seg, "openable_links_incomplete_reason")
        assert seg.openable_links_count == 0
        assert seg.openable_links_status == "none"
        assert seg.openable_links_incomplete_reason == ""


# ---------------------------------------------------------------------------
# assemble_correction 可開啟連結計算
# ---------------------------------------------------------------------------

class TestAssembleOpenableLinks:
    """assemble_correction 正確計算可開啟連結數量與狀態。"""

    def test_sufficient_with_two_cited_urls(self):
        """兩個引用來源都有有效 URL → sufficient。"""
        src_a = _source("s1", "來源A", url="https://a.com/1")
        src_b = _source("s2", "來源B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_count >= MIN_OPENABLE_LINKS
        assert seg.openable_links_status == "sufficient"
        assert seg.openable_links_incomplete_reason == ""

    def test_insufficient_with_one_cited_url(self):
        """只有一個引用來源有有效 URL → insufficient。"""
        src = _source("s1", "來源A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_count == 1
        assert seg.openable_links_status == "insufficient"
        assert "不足" in seg.openable_links_incomplete_reason

    def test_insufficient_with_no_valid_urls(self):
        """引用來源 URL 全為 None → insufficient。"""
        src = _source("s1", "來源A", url=None)
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_count == 0
        assert seg.openable_links_status == "insufficient"
        assert "無可開啟連結" in seg.openable_links_incomplete_reason

    def test_sufficient_with_extended_readings(self):
        """引用來源不足但延伸閱讀補足 → sufficient。"""
        src_cited = _source("s1", "引用來源", url="https://a.com/1")
        src_ext = _source("s2", "延伸來源", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources(
            [src_cited, src_ext], gap,
            used_ids=["s1"], omitted_ids=["s2"],
        )
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_count >= MIN_OPENABLE_LINKS
        assert seg.openable_links_status == "sufficient"

    def test_no_candidates_pure_pending(self):
        """無任何候選來源（pure pending）→ sufficient（不強制要求）。"""
        gap = Gap("問題？", "missing", "未說明")
        doc = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: WrittenSupplement("【待補證】尚無可用來源。", [])},
            {gap.question: cross_validate(gap.question, [])},
        )
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_status == "sufficient"
        assert seg.openable_links_incomplete_reason == ""

    def test_mixed_cited_and_extended_urls(self):
        """1 條引用 URL + 1 條延伸閱讀 URL → sufficient。"""
        src_cited = _source("s1", "引用", url="https://a.com/1")
        src_ext = _source("s2", "延伸", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources(
            [src_cited, src_ext], gap,
            used_ids=["s1"], omitted_ids=["s2"],
        )
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        assert seg.openable_links_count >= 2
        assert seg.openable_links_status == "sufficient"


def _assemble_with_sources(
    sources: list[Source],
    gap: Gap,
    used_ids: list[str] | None = None,
    omitted_ids: list[str] | None = None,
) -> CorrectionDoc:
    """組裝含指定來源的 CorrectionDoc。"""
    used_ids = used_ids or [s.id for s in sources]
    text = "補充內容[^1]。" if len(used_ids) >= 1 else "【待補證】"
    return assemble_correction(
        _doc(),
        [gap],
        {gap.question: sources},
        {gap.question: WrittenSupplement(text, used_ids, omitted_source_ids=omitted_ids or [])},
        {gap.question: cross_validate(gap.question, sources)},
    )


# ---------------------------------------------------------------------------
# binding_report 可開啟連結欄位
# ---------------------------------------------------------------------------

class TestBindingReportOpenableLinks:
    """binding_report 必須包含可開啟連結欄位與 checks。"""

    def test_report_has_fields(self):
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        for arg in report["arguments"]:
            assert "openable_links_count" in arg
            assert "openable_links_status" in arg
            assert "openable_links_incomplete_reason" in arg
            assert "at_least_two_openable_links" in arg["checks"]

    def test_parse_validates_fields(self):
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        parse_binding_report(report)  # 不 raise

    def test_check_matches_count(self):
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        report = build_binding_report(doc)
        arg = report["arguments"][0]
        expected = arg["openable_links_count"] >= MIN_OPENABLE_LINKS or arg["openable_links_status"] == "sufficient"
        assert arg["checks"]["at_least_two_openable_links"] is expected

    def test_insufficient_status_must_have_reason(self):
        """insufficient 狀態必須有 incomplete_reason。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        report = build_binding_report(doc)
        arg = report["arguments"][0]
        if arg["openable_links_status"] == "insufficient":
            assert arg["openable_links_incomplete_reason"].strip()

    def test_sufficient_status_no_reason(self):
        """sufficient 狀態不可有 incomplete_reason。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        arg = report["arguments"][0]
        if arg["openable_links_status"] == "sufficient":
            assert not arg["openable_links_incomplete_reason"].strip()


# ---------------------------------------------------------------------------
# binding_report 負例
# ---------------------------------------------------------------------------

class TestOpenableLinksNegativeCases:
    """負例：欄位缺失或不一致時 parse_binding_report 應明確失敗。"""

    def test_missing_openable_links_count_rejected(self):
        """缺少 openable_links_count → raise。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        del report["arguments"][0]["openable_links_count"]
        with pytest.raises(ValueError, match="缺少欄位"):
            parse_binding_report(report)

    def test_invalid_openable_links_status_rejected(self):
        """openable_links_status 非法值 → raise。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        report["arguments"][0]["openable_links_status"] = "invalid"
        with pytest.raises(ValueError, match="openable_links_status 非法"):
            parse_binding_report(report)

    def test_negative_count_rejected(self):
        """openable_links_count 為負數 → raise。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        report["arguments"][0]["openable_links_count"] = -1
        with pytest.raises(ValueError, match="openable_links_count 必須為非負整數"):
            parse_binding_report(report)

    def test_check_count_mismatch_rejected(self):
        """checks.at_least_two_openable_links 與實際數量不一致 → raise。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        report["arguments"][0]["checks"]["at_least_two_openable_links"] = False
        with pytest.raises(ValueError, match="at_least_two_openable_links 與 openable_links_count 不一致"):
            parse_binding_report(report)

    def test_sufficient_with_reason_rejected(self):
        """sufficient 但有 incomplete_reason → raise。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        report = build_binding_report(doc)
        report["arguments"][0]["openable_links_incomplete_reason"] = "不應出現"
        with pytest.raises(ValueError, match="sufficient 但 incomplete_reason 非空"):
            parse_binding_report(report)

    def test_insufficient_without_reason_rejected(self):
        """insufficient 但無 incomplete_reason → raise。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        report = build_binding_report(doc)
        report["arguments"][0]["openable_links_incomplete_reason"] = ""
        with pytest.raises(ValueError, match="insufficient 但 incomplete_reason 為空"):
            parse_binding_report(report)


# ---------------------------------------------------------------------------
# export 輸出：【待補來源】標記
# ---------------------------------------------------------------------------

class TestOpenableLinksInExport:
    """JSON/Markdown/DOCX 輸出必須包含可開啟連結欄位與標記。"""

    def test_json_has_fields(self):
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        data = to_json(doc)
        seg = [s for s in data["segments"] if s["type"] == "supplement"][0]
        assert "openable_links_count" in seg
        assert "openable_links_status" in seg
        assert "openable_links_incomplete_reason" in seg

    def test_markdown_insufficient_marker(self):
        """Markdown 輸出 insufficient 時必須有【待補來源】標記。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        md = to_markdown(doc)
        assert "【待補來源】" in md
        assert "可開啟連結不足" in md

    def test_markdown_sufficient_no_marker(self):
        """Markdown 輸出 sufficient 時不應有【待補來源】標記。"""
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        md = to_markdown(doc)
        assert "可開啟連結不足" not in md

    def test_docx_insufficient_marker(self):
        """DOCX 輸出 insufficient 時必須有【待補來源】段落。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        to_docx(doc, out_path)
        out_doc = DocxDocument(out_path)
        texts = [p.text for p in out_doc.paragraphs]
        assert any("【待補來源】" in t and "可開啟連結不足" in t for t in texts)

    def test_original_text_unchanged(self):
        """可開啟連結欄位不影響原稿逐字不變。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        data = to_json(doc)
        originals = [s for s in data["segments"] if s["type"] == "original"]
        for seg in originals:
            assert seg["text"] == "原稿逐字保留。"


# ---------------------------------------------------------------------------
# 禁止拼湊 URL / 混排假結果
# ---------------------------------------------------------------------------

class TestNoFabricatedUrls:
    """禁止拼湊 URL、禁止把待補來源與真實連結混排。"""

    def test_no_urls_in_output_when_sources_lack_urls(self):
        """來源無 URL 時，不應在成品中產生虛構 URL。"""
        src = _source("s1", "A", url=None)
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        data = to_json(doc)
        seg = [s for s in data["segments"] if s["type"] == "supplement"][0]
        # sources 中不應有有效 URL
        for s in seg["sources"]:
            assert s["url"] is None or not is_openable_url(s["url"])

    def test_extended_readings_only_from_retrieved(self):
        """延伸閱讀只含實際檢索到的候選來源，不可虛構。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        seg = [s for s in doc.segments if s.type == "supplement"][0]
        # 無 omitted sources → extended_readings 應為空
        assert seg.extended_readings == []

    def test_binding_report_no_fabricated_source_ids(self):
        """binding_report 中 source_ids 只含實際引用來源。"""
        src = _source("s1", "A", url="https://a.com/1")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src], gap, used_ids=["s1"])
        report = build_binding_report(doc)
        arg = report["arguments"][0]
        assert arg["source_ids"] == ["s1"]


# ---------------------------------------------------------------------------
# 原稿不可變性
# ---------------------------------------------------------------------------

class TestOriginalImmutability:
    """可開啟連結欄位不影響原稿逐字不變。"""

    def test_original_segment_unchanged(self):
        src_a = _source("s1", "A", url="https://a.com/1")
        src_b = _source("s2", "B", url="https://b.com/2")
        gap = Gap("問題？", "missing", "未說明")
        doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
        original = [s for s in doc.segments if s.type == "original"]
        assert len(original) == 1
        assert original[0].text == "原稿逐字保留。"
