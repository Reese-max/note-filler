"""來源附註負例：證據不足或適用條件未明時不得生成確定結論。"""

import json
from urllib.parse import urlparse

from docx import Document as DocxDocument

from note_filler.binding_report import build_binding_report, parse_binding_report
from note_filler.correction import assemble_correction
from note_filler.export import to_docx, to_json, to_markdown
from note_filler.gap import Gap
from note_filler.parse import Document, Paragraph
from note_filler.retrieve.models import Source
from note_filler.verify import cross_validate
from note_filler.write import WrittenSupplement


_ORIGINAL = "原稿逐字保留。"


def _source(
    source_id: str,
    title: str,
    content: str,
    level: str,
    url: str | None = None,
) -> Source:
    return Source(
        id=source_id,
        title=title,
        url=url,
        level=level,
        content=content,
        fetched_date="2026-07-28",
        doc_date=None,
        distance=0.1,
    )


def _product(sources: list[Source], text: str):
    question = "此申請是否符合要件？"
    gap = Gap(question, "missing", "原稿未說明申請要件")
    validation = cross_validate(question, sources)
    product = assemble_correction(
        Document("input/note.txt", (Paragraph(0, _ORIGINAL),), _ORIGINAL),
        [gap],
        {question: sources},
        {question: WrittenSupplement(text, [source.id for source in sources])},
        {question: validation},
    )
    return product, validation


def _exports(product, tmp_path, name: str):
    data = to_json(product)
    markdown = to_markdown(product)
    path = tmp_path / f"{name}.docx"
    to_docx(product, str(path))
    paragraphs = [paragraph.text for paragraph in DocxDocument(path).paragraphs]
    supplement = next(
        segment for segment in data["segments"] if segment["type"] == "supplement"
    )
    assert next(
        segment["text"] for segment in data["segments"] if segment["type"] == "original"
    ) == _ORIGINAL
    assert _ORIGINAL in markdown and paragraphs[0] == _ORIGINAL
    return supplement, markdown, paragraphs


def _assert_conflict_is_automatically_marked_pending(
    product, validation, expected_source_ids, tmp_path, name: str
):
    """驗證衝突訊號由流程傳遞，不以測試手動覆寫 confidence 偽造。"""
    assert validation.conflict is True
    segment = product.segments[-1]
    assert segment.confidence == "pending_evidence"
    assert segment.conflict_note == validation.conflict_note
    assert segment.source_ids == expected_source_ids
    assert segment.pending_evidence_reason == "引用來源衝突且無法判定適用條件"

    supplement, markdown, paragraphs = _exports(product, tmp_path, name)
    assert supplement["conflict_note"] == validation.conflict_note
    annotation = supplement["three_part_annotation"]
    assert [row["id"] for row in annotation["source_comparison"]] == expected_source_ids
    assert all("【待補來源】" in item for item in annotation["usage_conditions"])
    assert annotation["conclusion"].startswith("【待補來源】")
    assert "不合併為單一結論" in annotation["conclusion"]
    assert "【待補來源】" in markdown
    assert any("結論：【待補來源】" in line for line in paragraphs)

    report = parse_binding_report(build_binding_report(product))
    argument = report["arguments"][0]
    conflict_fields = (
        "source_conflicts",
        "source_preference_reason",
        "applicable_conditions",
        "readable_conclusion",
    )
    assert all(field in argument for field in conflict_fields)
    assert argument["source_conflicts"][0]["source_ids"] == expected_source_ids
    assert argument["source_preference_reason"]["status"] == "not_selected"
    assert argument["applicable_conditions"] == {"status": "not_assessed", "items": []}
    assert argument["readable_conclusion"] == (
        "來源表述不一致；尚未選定優先來源，需人工判讀適用條件。"
    )


def test_unreferenced_conflict_candidates_do_not_degrade_referenced_argument(tmp_path):
    question = "此申請是否符合要件？"
    cited_primary = _source(
        "official:primary",
        "申請文件審核紀錄",
        "本案申請表件及身分資料均已備齊。",
        "A",
        "https://example.test/primary",
    )
    cited_secondary = _source(
        "official:secondary",
        "申請程序受理紀錄",
        "主管機關已受理本案申請文件。",
        "A",
        "https://example.test/secondary",
    )
    unreferenced_allow = _source(
        "catalog:allow",
        "候選來源甲",
        "申請人得提出申請。",
        "A",
    )
    unreferenced_deny = _source(
        "catalog:deny",
        "候選來源乙",
        "申請人不得提出申請。",
        "A",
    )
    candidates = [
        cited_primary,
        cited_secondary,
        unreferenced_allow,
        unreferenced_deny,
    ]
    validation = cross_validate(question, candidates)
    assert validation.conflict is True

    product = assemble_correction(
        Document("input/note.txt", (Paragraph(0, _ORIGINAL),), _ORIGINAL),
        [Gap(question, "missing", "原稿未說明申請要件")],
        {question: candidates},
        {
            question: WrittenSupplement(
                "本案已備齊申請要件。[^1][^2]",
                [cited_primary.id, cited_secondary.id],
            )
        },
        {question: validation},
    )

    segment = product.segments[-1]
    assert segment.source_ids == [cited_primary.id, cited_secondary.id]
    assert segment.confidence == "verified"
    assert segment.conflict_note is None

    supplement, markdown, _ = _exports(product, tmp_path, "unreferenced-conflict")
    assert supplement["three_part_annotation"]["conclusion"].startswith("建議以")
    assert "【待補來源】" not in markdown
    assert "衝突告警" not in markdown

    argument = parse_binding_report(build_binding_report(product))["arguments"][0]
    assert argument["source_conflicts"] == []
    assert argument["source_preference_reason"]["reason_code"] == "no_conflict"


def test_claim_without_qualified_source_outputs_pending_without_fabricated_url(tmp_path):
    candidate = _source(
        "commentary:1",
        "未經查核的單一評論",
        "這是一般性評論，沒有可獨立驗證的依據。",
        "D",
    )
    product, validation = _product([candidate], "此申請應予准許。[^1]")

    assert validation.verified is False
    assert product.segments[-1].confidence == "pending_evidence"
    supplement, markdown, paragraphs = _exports(product, tmp_path, "unqualified")
    annotation = supplement["three_part_annotation"]

    assert [row["url"] for row in annotation["source_comparison"]] == [None]
    assert all("【待補來源】" in item for item in annotation["usage_conditions"])
    assert annotation["conclusion"].startswith("【待補來源】")
    assert "建議以" not in annotation["conclusion"]
    assert "http://" not in markdown and "https://" not in markdown
    assert not any("http://" in line or "https://" in line for line in paragraphs)


def test_claim_with_unattributable_source_difference_is_automatically_marked_pending(
    tmp_path,
):
    active = _source(
        "catalog:current",
        "資料庫甲的現行摘錄",
        "資料庫甲將同名審查基準標示為有效，但未附版本、發布日或適用機關。",
        "A",
    )
    repealed = _source(
        "catalog:archived",
        "資料庫乙的歷史摘錄",
        "資料庫乙將同名審查基準標示為廢止，但未附版本、廢止日或適用機關。",
        "A",
    )
    product, validation = _product(
        [active, repealed], "本案審查基準是否仍可適用？[^1][^2]"
    )

    _assert_conflict_is_automatically_marked_pending(
        product,
        validation,
        [active.id, repealed.id],
        tmp_path,
        "unattributable-source-difference",
    )


def test_sources_with_unknown_applicability_are_automatically_marked_pending(
    tmp_path,
):
    ordinary = _source(
        "procedure:ordinary",
        "一般程序摘錄",
        "來源甲認定一般程序的核定結果合法，但未記載程序類型、主體或適用期間。",
        "A",
    )
    special = _source(
        "procedure:special",
        "特別程序摘錄",
        "來源乙認定特別程序的核定結果違法，但未記載程序類型、主體或適用期間。",
        "A",
    )
    product, validation = _product(
        [ordinary, special], "本案應採何種程序並如何評估核定結果？[^1][^2]"
    )

    _assert_conflict_is_automatically_marked_pending(
        product,
        validation,
        [ordinary.id, special.id],
        tmp_path,
        "unknown-applicability",
    )


def _is_http_url(value: object) -> bool:
    if not isinstance(value, str):
        return False
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def test_each_argument_block_has_parseable_source_links_and_conflict_summary():
    """最終 JSON 的每個論點逐筆鎖定來源連結、待補互斥與衝突歸屬。"""
    verified = _source(
        "law:verified",
        "行政程序法第 1 條",
        "行政程序應依法進行。",
        "A",
        "https://law.moj.gov.tw/LawClass/LawSingle.aspx?pcode=A0030055&flno=1",
    )
    conflict_active = _source(
        "law:current",
        "資料庫甲的現行摘錄",
        "資料庫甲將同名審查基準標示為有效，但未附版本、發布日或適用機關。",
        "A",
    )
    conflict_repealed = _source(
        "law:archived",
        "資料庫乙的歷史摘錄",
        "資料庫乙將同名審查基準標示為廢止，但未附版本、廢止日或適用機關。",
        "A",
    )
    questions = ["程序依據為何？", "尚缺何項資料？", "審查基準是否仍可適用？"]
    gaps = [
        Gap(questions[0], "missing", "原稿未說明程序依據"),
        Gap(questions[1], "missing", "原稿未列出待補資料"),
        Gap(questions[2], "missing", "原稿未說明申請資格"),
    ]
    conflict = cross_validate(questions[2], [conflict_active, conflict_repealed])
    product = assemble_correction(
        Document("input/note.txt", (Paragraph(0, _ORIGINAL),), _ORIGINAL),
        gaps,
        {
            questions[0]: [verified],
            questions[1]: [],
            questions[2]: [conflict_active, conflict_repealed],
        },
        {
            questions[0]: WrittenSupplement("程序應依法進行。[^1]", [verified.id]),
            questions[1]: WrittenSupplement("【待補證】尚無可用來源。", []),
            questions[2]: WrittenSupplement(
                "現有資料對審查基準效力互有差異。[^1][^2]",
                [conflict_active.id, conflict_repealed.id],
            ),
        },
        {
            questions[0]: cross_validate(questions[0], [verified]),
            questions[1]: cross_validate(questions[1], []),
            questions[2]: conflict,
        },
    )

    data = json.loads(json.dumps(to_json(product), ensure_ascii=False))
    arguments = [segment for segment in data["segments"] if segment["type"] == "supplement"]

    assert [argument["argument_id"] for argument in arguments] == [
        "argument:0",
        "argument:1",
        "argument:2",
    ]

    conflict_summaries: dict[str, str] = {}
    for argument in arguments:
        annotation = argument["three_part_annotation"]
        rows = annotation["source_comparison"]
        links = [row["url"] for row in rows if row["url"]]
        has_pending_source = "【待補來源】" in json.dumps(
            annotation, ensure_ascii=False
        )

        assert len(links) == sum(source["url"] is not None for source in argument["sources"])
        assert all(_is_http_url(link) for link in links)
        assert not (has_pending_source and links)

        parsed_conflicts = [
            note.removeprefix("來源衝突：")
            for note in annotation["discrepancy_notes"]
            if note.startswith("來源衝突：")
        ]
        if argument["conflict_note"]:
            assert parsed_conflicts == [argument["conflict_note"]]
            conflict_summaries[argument["argument_id"]] = parsed_conflicts[0]
        else:
            assert parsed_conflicts == []

    assert conflict.conflict is True
    assert conflict_summaries == {"argument:2": conflict.conflict_note}
