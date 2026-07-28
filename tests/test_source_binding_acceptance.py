"""驗收測試：生成結果中每個論點對應到至少一個來源，且來源集合不重複遺漏。
涵蓋一對一與一對多綁定情境。

直接驗證 assemble_correction 與 write_supplement 形成的
Segment.sources / Segment.traceability / Segment.source_id 之間的不變式。

每個 supplement segment 的來源不變式：
  (a) 完整性：所有 used_source_ids 都有對應 Source 物件在 segment.sources
  (b) 無重複：segment.sources 內無重複 source.id
  (c) 無多餘：segment.sources 中的來源皆有在 used_source_ids 中
  (d) 追溯一致性：traceability 與 sources 一一對應
"""
from __future__ import annotations

import re

import pytest

from note_filler.binding_report import (
    build_binding_report,
    parse_binding_report,
    write_binding_report,
)
from note_filler.correction import assemble_correction
from note_filler.gap import Gap
from note_filler.parse import Document, Paragraph
from note_filler.retrieve.models import Source
from note_filler.verify import cross_validate
from note_filler.write import WrittenSupplement


# ---- 輔助函式 ---------------------------------------------------------------

def _doc():
    return Document(
        source_path="input/note.txt",
        paragraphs=(Paragraph(0, "原稿逐字保留。"),),
        full_text="原稿逐字保留。",
    )


def _source(id: str, title: str, level: str = "A") -> Source:
    return Source(
        id=id, title=title, url=None, level=level,
        content=f"{title} 內容", fetched_date="2026-07-24",
        doc_date=None, distance=0.5,
    )


def _assert_source_binding_invariants(segments) -> None:
    """對每個 supplement segment 驗證來源綁定不變式。"""
    for i, seg in enumerate(segments):
        if seg.type != "supplement":
            continue
        if not seg.sources:
            # 無來源段不需驗證綁定
            continue

        # (a) 無重複 id
        ids = [s.id for s in seg.sources]
        assert len(ids) == len(set(ids)), (
            f"segment[{i}] sources 含重複 id: {ids}"
        )

        # (b) traceability 與 sources 一一對應
        expected_trace = [{"kind": "source", "id": s.id} for s in seg.sources]
        assert seg.traceability == expected_trace, (
            f"segment[{i}] traceability 與 sources 不對應: "
            f"trace={[[t['kind'], t['id']] for t in seg.traceability]}, "
            f"expected={expected_trace}"
        )

        # (c) source_id 格式與內容正確
        expected_sid = f"sources:{','.join(ids)}"
        assert seg.source_id == expected_sid, (
            f"segment[{i}] source_id 應為 {expected_sid!r}, 實際 {seg.source_id!r}"
        )

        # (d) 所有 source 都有合法 level
        assert all(s.level in ("A", "B", "C", "D") for s in seg.sources), (
            f"segment[{i}] 含不合法的 level"
        )


def _assert_used_ids_covered(segment, used_ids: list[str]) -> None:
    """斷言 used_source_ids 中的每個 id 都在 segment.sources 且有對應 traceability。"""
    source_ids = {s.id for s in segment.sources}
    trace_ids = {t["id"] for t in segment.traceability if t.get("kind") == "source"}
    for uid in used_ids:
        assert uid in source_ids, (
            f"used_source_id {uid!r} 不在 segment.sources 中"
        )
        assert uid in trace_ids, (
            f"used_source_id {uid!r} 不在 segment.traceability 中"
        )


def _assert_no_extra_sources(segment, used_ids: list[str]) -> None:
    """斷言 segment.sources 中沒有未在 used_source_ids 出現的來源。"""
    for s in segment.sources:
        assert s.id in used_ids, (
            f"segment.sources 中有未使用的來源: {s.id!r}"
        )


# ---- 一對一綁定：每個 supplement 段對應剛好一個來源 -----------------------

def test_one_to_one_single_source():
    """一個 gap、一個來源、一個 [^1] 標記 → 1:1 綁定。"""
    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("行政處分之定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("定義參照[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    assert seg.type == "supplement"
    assert len(seg.sources) == 1
    assert seg.sources[0].id == src.id
    assert len(seg.traceability) == 1
    assert seg.traceability[0] == {"kind": "source", "id": src.id}
    assert seg.source_id == "sources:law:92"
    _assert_source_binding_invariants(product.segments)


def test_one_to_one_two_separate_gaps():
    """兩個獨立 gap，各自綁定一個來源 → 每段都是 1:1。"""
    src_a = _source("law:92", "行政程序法第 92 條")
    src_b = _source("law:93", "行政程序法第 93 條")
    gap_a = Gap("定義？", "missing", "未說明")
    gap_b = Gap("附款？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b]},
        {
            gap_a.question: WrittenSupplement("定義[^1]。", [src_a.id]),
            gap_b.question: WrittenSupplement("附款[^1]。", [src_b.id]),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b]),
        },
    )
    assert len(product.segments) == 3  # 1 original + 2 supplement
    seg_a, seg_b = product.segments[1], product.segments[2]
    assert seg_a.type == seg_b.type == "supplement"
    assert len(seg_a.sources) == len(seg_b.sources) == 1
    assert seg_a.sources[0].id == src_a.id
    assert seg_b.sources[0].id == src_b.id
    assert seg_a.source_id == "sources:law:92"
    assert seg_b.source_id == "sources:law:93"
    _assert_source_binding_invariants(product.segments)


# ---- 一對多綁定：一個 supplement 段引用多個來源 ---------------------------

def test_one_to_many_three_sources():
    """一個 gap 引用三個來源 → 1:3 綁定。"""
    srcs = [
        _source("law:92", "行政程序法第 92 條"),
        _source("law:93", "行政程序法第 93 條", level="B"),
        _source("web:abc", "官方函釋", level="C"),
    ]
    gap = Gap("附款相關規定？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: srcs},
        {gap.question: WrittenSupplement("三源[^1][^2][^3]。", [s.id for s in srcs])},
        {gap.question: cross_validate(gap.question, srcs)},
    )
    seg = product.segments[1]
    assert seg.type == "supplement"
    assert len(seg.sources) == 3
    assert [s.id for s in seg.sources] == ["law:92", "law:93", "web:abc"]
    assert len(seg.traceability) == 3
    assert seg.traceability == [
        {"kind": "source", "id": "law:92"},
        {"kind": "source", "id": "law:93"},
        {"kind": "source", "id": "web:abc"},
    ]
    assert seg.source_id == "sources:law:92,law:93,web:abc"
    _assert_source_binding_invariants(product.segments)


def test_one_to_many_with_one_to_one_in_same_doc():
    """混合情境：gapA 一對一、gapB 一對多。"""
    src_a = _source("law:92", "行政程序法第 92 條")
    src_b1 = _source("law:93", "行政程序法第 93 條")
    src_b2 = _source("law:94", "行政程序法第 94 條")
    gap_a = Gap("定義？", "missing", "未說明")
    gap_b = Gap("附款限制？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b1, src_b2]},
        {
            gap_a.question: WrittenSupplement("定義[^1]。", [src_a.id]),
            gap_b.question: WrittenSupplement("限制[^1][^2]。", [src_b1.id, src_b2.id]),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b1, src_b2]),
        },
    )
    assert len(product.segments) == 3
    seg_a, seg_b = product.segments[1], product.segments[2]
    assert seg_a.source_id == "sources:law:92"
    assert seg_b.source_id == "sources:law:93,law:94"
    assert len(seg_a.sources) == 1
    assert len(seg_b.sources) == 2
    _assert_source_binding_invariants(product.segments)


# ---- 來源集合完整性：無遺漏 ------------------------------------------------

def test_no_omitted_source_all_used_ids_have_source_objects():
    """所有 used_source_ids 都在 segment.sources 中有對應 Source 物件。"""
    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    _assert_used_ids_covered(product.segments[1], [src.id])


def test_no_omitted_source_multiple_ids():
    """多個 used_source_ids 全數在 segment.sources 中有對應。"""
    srcs = [
        _source("law:92", "行政程序法第 92 條"),
        _source("law:93", "行政程序法第 93 條", level="B"),
        _source("law:94", "行政程序法第 94 條", level="B"),
    ]
    used = [s.id for s in srcs]
    gap = Gap("規定？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: srcs},
        {gap.question: WrittenSupplement("三源[^1][^2][^3]。", used)},
        {gap.question: cross_validate(gap.question, srcs)},
    )
    _assert_used_ids_covered(product.segments[1], used)


def test_partial_usage_no_omitted():
    """檢索了五個來源但只引用兩個 → 只 used 的兩筆有對應。"""
    srcs = [
        _source("law:92", "行政程序法第 92 條"),
        _source("law:93", "行政程序法第 93 條"),
        _source("law:94", "行政程序法第 94 條"),
        _source("web:abc", "官方函釋", level="C"),
        _source("web:xyz", "實務見解", level="D"),
    ]
    used_ids = ["law:92", "law:94"]
    used_sources = [s for s in srcs if s.id in used_ids]
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: srcs},
        {gap.question: WrittenSupplement("兩源[^1][^2]。", used_ids)},
        {gap.question: cross_validate(gap.question, used_sources)},
    )
    seg = product.segments[1]
    assert len(seg.sources) == 2
    assert seg.sources[0].id == "law:92"
    assert seg.sources[1].id == "law:94"
    _assert_used_ids_covered(seg, used_ids)
    _assert_no_extra_sources(seg, used_ids)


# ---- 來源集合無多餘 ---------------------------------------------------------

def test_no_extra_sources_not_in_used():
    """segment.sources 不包含未被使用的來源 id。"""
    srcs = [
        _source("law:92", "行政程序法第 92 條"),
        _source("law:93", "行政程序法第 93 條"),
    ]
    used_ids = ["law:92"]  # 只引用了第一個
    used_sources = [s for s in srcs if s.id in used_ids]
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: srcs},
        {gap.question: WrittenSupplement("定義[^1]。", used_ids)},
        {gap.question: cross_validate(gap.question, used_sources)},
    )
    seg = product.segments[1]
    assert len(seg.sources) == 1
    assert seg.sources[0].id == "law:92"
    _assert_no_extra_sources(seg, used_ids)


# ---- 來源集合無重複 ---------------------------------------------------------

def test_no_duplicate_sources():
    """segment.sources 無重複 source.id。"""
    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: [src, src]},  # retrieved 有兩個相同 id 的 Source
        {gap.question: WrittenSupplement("定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    ids = [s.id for s in seg.sources]
    assert len(ids) == len(set(ids)), f"sources 含重複 id: {ids}"
    assert ids == [src.id]  # 應只保留一個


def test_no_duplicate_traceability():
    """segment.traceability 無重複條目。"""
    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    trace_ids = [t["id"] for t in seg.traceability]
    assert len(trace_ids) == len(set(trace_ids)), f"traceability 含重複 id: {trace_ids}"


# ---- 綁定不變式 x 多個 segment -----------------------------------------------

def test_all_supplement_segments_have_valid_binding():
    """所有 supplement segment 都滿足來源綁定不變式。"""
    srcs_a = [_source("law:92", "行政程序法第 92 條")]
    srcs_b = [
        _source("law:93", "行政程序法第 93 條"),
        _source("law:94", "行政程序法第 94 條", level="B"),
    ]
    gap_a = Gap("定義？", "missing", "未說明")
    gap_b = Gap("附款？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap_a, gap_b],
        {gap_a.question: srcs_a, gap_b.question: srcs_b},
        {
            gap_a.question: WrittenSupplement("定義[^1]。", [s.id for s in srcs_a]),
            gap_b.question: WrittenSupplement("附款限制[^1][^2]。", [s.id for s in srcs_b]),
        },
        {
            gap_a.question: cross_validate(gap_a.question, srcs_a),
            gap_b.question: cross_validate(gap_b.question, srcs_b),
        },
    )
    _assert_source_binding_invariants(product.segments)


def test_source_binding_survives_require_traceable_gate():
    """綁定不變式驗證後再通過 require_traceable_note_product 閘。"""
    from note_filler.pipeline import require_traceable_note_product
    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("定義？", "missing", "未說明")
    product = assemble_correction(
        _doc(), [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    _assert_source_binding_invariants(product.segments)
    require_traceable_note_product(product, source="binding-acceptance")


# ---- 失敗語義：論點存在但來源缺失／來源未對上（最小負例） -------------------
#
# 鎖定：不得默默輸出「有論點、無／錯綁定」的不完整成品。
# assemble 可能先留下稽核與不一致 state；末端 require_traceable 必須硬失敗，
# 並在錯誤訊息／稽核中指出缺少或未對上的綁定。


def test_failure_semantic_claim_exists_but_source_missing(caplog):
    """最小負例：補充論點文字存在，但 used_source_ids 在 retrieved 完全找不到。

    必須：
    1) 稽核 used_sources_not_forwarded 列出 missing_source_ids
    2) require_traceable_note_product 明確 raise（不得默默過關）
    3) 錯誤語意指向綁定／source_id 缺口
    """
    import json
    import logging

    from note_filler.export import to_markdown
    from note_filler.pipeline import require_traceable_note_product

    gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
    claim = "行政處分應符合法定要件，並保障當事人陳述意見之機會。"
    missing_id = "missing-src-BOUND-01"
    with caplog.at_level(logging.WARNING):
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},  # 來源集合為空
            {gap.question: WrittenSupplement(claim, [missing_id])},
            {},
        )
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗") as ei:
            require_traceable_note_product(product, source="fail-bind-missing")
    seg = product.segments[-1]
    assert seg.type == "supplement"
    assert claim in seg.text  # 論點存在
    assert seg.sources == []  # 來源缺失
    assert missing_id in seg.source_id  # 宣稱綁定了不存在的來源 → 不一致

    # 若略過閘門，markdown 會帶論點卻無掛上 Source 物件 → 不完整輸出風險
    md = to_markdown(product)
    assert claim in md

    msg = str(ei.value)
    assert "source_id" in msg or "pending:gap:" in msg or "對應失敗" in msg, (
        f"失敗訊息應指出綁定缺口，實際：{msg!r}"
    )
    assert "note_traceability_failed" in caplog.text
    assert "fail-bind-missing" in caplog.text

    # 組裝階段即應留下缺失來源稽核（可機器比對）
    audit_events = []
    for rec in caplog.records:
        try:
            payload = json.loads(rec.message)
        except (json.JSONDecodeError, TypeError):
            continue
        if payload.get("event") == "used_sources_not_forwarded":
            audit_events.append(payload)
    assert audit_events, "應有 used_sources_not_forwarded 稽核"
    assert any(
        missing_id in json.dumps(p.get("missing_source_ids", []), ensure_ascii=False)
        for p in audit_events
    ), f"稽核應列出缺失綁定 id {missing_id!r}"


def test_failure_semantic_claim_exists_but_source_not_aligned(caplog):
    """最小負例：論點引用多個來源 id，但 retrieved 只對上一部分 → 來源未對上。

    source_id 宣稱 a,b-missing，實際 sources 只有 a → 閘必須硬失敗並指出 source_id 不符。
    """
    import json
    import logging

    from note_filler.pipeline import require_traceable_note_product

    gap = Gap("附款之限制為何？", "missing", "原稿未列")
    present = _source("law:92", "行政程序法第 92 條")
    missing_id = "law:ghost-BOUND-02"
    claim = "附款不得違背行政處分之目的[^1][^2]。"
    with caplog.at_level(logging.WARNING):
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [present]},
            {gap.question: WrittenSupplement(claim, [present.id, missing_id])},
            {gap.question: cross_validate(gap.question, [present])},
        )
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗|source_id") as ei:
            require_traceable_note_product(product, source="fail-bind-unaligned")
    seg = product.segments[-1]
    assert claim.split("[")[0] in seg.text or "附款" in seg.text
    assert [s.id for s in seg.sources] == [present.id]
    assert seg.source_id == f"sources:{present.id},{missing_id}"
    assert seg.traceability == [{"kind": "source", "id": present.id}]

    msg = str(ei.value)
    assert "source_id" in msg, f"應指出 source_id 未對上，實際：{msg!r}"
    assert present.id in msg
    assert "note_traceability_failed" in caplog.text

    missing_audits = []
    for rec in caplog.records:
        try:
            payload = json.loads(rec.message)
        except (json.JSONDecodeError, TypeError):
            continue
        if payload.get("event") == "used_sources_not_forwarded":
            missing_audits.append(payload)
    assert any(
        missing_id in json.dumps(p.get("missing_source_ids", []), ensure_ascii=False)
        for p in missing_audits
    ), f"應回報缺失綁定 {missing_id!r}"


def test_failure_semantic_source_trace_not_matched_to_sources(caplog):
    """最小負例：論點已掛 Source 物件，但 traceability 指到另一 id → 來源未對上。

    模擬綁定漂移；閘必須 raise「source ID 對應失敗」，不得匯出不一致成品。
    """
    import logging

    from note_filler.export import to_markdown
    from note_filler.pipeline import require_traceable_note_product

    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("定義？", "missing", "未說明")
    claim = "行政處分之定義參照法定要件[^1]。"
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement(claim, [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    # 人為製造「來源未對上」：trace 指到不存在的綁定
    product.segments[-1].traceability = [{"kind": "source", "id": "wrong-bound-id"}]
    assert product.segments[-1].sources[0].id == src.id
    assert "定義" in product.segments[-1].text or "行政處分" in product.segments[-1].text

    # 證明若無閘門，不完整／不一致內容仍可序列化
    md = to_markdown(product)
    assert "行政處分" in md or "定義" in md

    with caplog.at_level(logging.WARNING):
        with pytest.raises(RuntimeError, match="source ID 對應失敗") as ei:
            require_traceable_note_product(product, source="fail-bind-trace")
    assert "來源追溯驗證失敗" in str(ei.value)
    assert "note_traceability_failed" in caplog.text
    assert "fail-bind-trace" in caplog.text


def test_failure_semantic_binding_report_flags_missing_source(caplog):
    """最小負例：assemble_correction 產出 source_id 宣稱存在但 sources 為空的 segment。

    驗證：
    1) 綁定報告的 binding_status 為 fail（不得默默標 pass）
    2) 錯誤訊息明確指出缺失綁定（含 source_id）
    3) 與 require_traceable_note_product 閘一致攔截
    """
    import logging

    from note_filler.pipeline import require_traceable_note_product

    gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
    claim = "行政處分應符合法定要件，並保障當事人陳述意見之機會。"
    missing_id = "missing-src-RPT-01"
    with caplog.at_level(logging.WARNING):
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},  # retrieved 為空：missing_id 在此找不到
            {gap.question: WrittenSupplement(claim, [missing_id])},
            {},
        )
    seg = product.segments[-1]
    assert seg.type == "supplement"
    assert claim in seg.text
    assert seg.sources == []
    assert missing_id in seg.source_id

    # 綁定報告必須明確標示 fail（不得靜默標 pass）
    report = parse_binding_report(build_binding_report(product))
    arg = report["arguments"][0]
    assert arg["binding_status"] == "fail", (
        f"缺失來源之 binding_status 應為 fail，實際：{arg['binding_status']!r}"
    )
    assert arg["binding_ok"] is False
    assert arg["checks"]["at_least_one_source"] is False
    assert report["summary"]["fail"] == 1
    assert report["summary"]["all_arguments_ok"] is False

    # 閘門必須明確 raise 且錯誤指向綁定缺口
    with caplog.at_level(logging.WARNING):
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗") as ei:
            require_traceable_note_product(product, source="fail-bind-rpt")
    msg = str(ei.value)
    assert "source_id" in msg or "對應失敗" in msg, (
        f"失敗訊息應指出綁定缺口，實際：{msg!r}"
    )
    assert "note_traceability_failed" in caplog.text


def test_failure_semantic_missing_binding_blocks_pipeline_return(tmp_path, monkeypatch, caplog):
    """主流程末端閘必須攔截缺綁定成品：run_pipeline 不得表面成功回傳。

    以 monkeypatch 在 assemble 後注入「有論點、無對上來源」的 segment，
    確認 RuntimeError 且無成功回傳（避免默默輸出不完整內容）。
    """
    import json
    import logging

    from docx import Document as DocxDocument

    from note_filler.correction import Segment
    from note_filler.llm import FakeLLM
    from note_filler import pipeline as pl
    from tests.test_pipeline import FakeLaw, FakeTwinkle

    note = tmp_path / "bind-fail.docx"
    d = DocxDocument()
    d.add_paragraph("行政程序法要求行政行為應遵守正當程序。")
    d.save(str(note))

    real_assemble = pl.assemble_correction

    def inject_broken_binding(*args, **kwargs):
        product = real_assemble(*args, **kwargs)
        # 附加：有實質論點、confidence 非 pending、無 sources／無合法 processing_record
        product.segments.append(
            Segment(
                type="supplement",
                text="此論點宣稱有據，但完全缺少來源綁定。",
                anchor_idx=None,
                sources=[],
                confidence="verified",  # 非 pending_evidence → 閘應判 processing_record 失敗
                traceability=[],
                source_id="",
                functional_gap="",
                user_value="",
                argument_id="",
            )
        )
        return product

    monkeypatch.setattr(pl, "assemble_correction", inject_broken_binding)

    llm = FakeLLM(
        [
            "admin",
            "正當程序的要件為何?",
            json.dumps(
                [
                    {
                        "question": "正當程序的要件為何?",
                        "status": "missing",
                        "reason": "未展開",
                    }
                ],
                ensure_ascii=False,
            ),
            '{"keyword": "正當程序", "law_name": null}',
            "【待補證】尚待補充。",
        ]
    )

    with caplog.at_level(logging.WARNING, logger="note_filler.pipeline"):
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗|processing_record|對應失敗"):
            pl.run_pipeline(str(note), llm, FakeTwinkle([[]]), FakeLaw())
    assert "note_traceability_failed" in caplog.text
    # 不得在失敗後仍以為成功（無回傳值可 assert；raise 即為攔截）
    assert not (tmp_path / "bind-fail.訂正稿.md").exists()


def test_failure_semantic_source_id_mismatch_vs_sources(caplog):
    """最小負例：論點有來源、source_ids 列表正確，但 source_id 欄位被竄改為不一致值。

    模擬序列化漂移或人為覆寫：source_id 應為 sources:law:92,law:93，
    實際卻寫成 sources:law:92,law:WRONG → 閘必須硬失敗並指出 source_id 未對上。
    """
    import logging

    from note_filler.pipeline import require_traceable_note_product

    src_a = _source("law:92", "行政程序法第 92 條")
    src_b = _source("law:93", "行政程序法第 93 條", level="B")
    gap = Gap("附款之限制為何？", "missing", "原稿未列")
    claim = "附款不得違背行政處分之目的[^1][^2]。"
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src_a, src_b]},
        {gap.question: WrittenSupplement(claim, [src_a.id, src_b.id])},
        {gap.question: cross_validate(gap.question, [src_a, src_b])},
    )
    seg = product.segments[-1]
    # 驗證組裝正確後，人為竄改 source_id 欄位
    assert seg.source_id == "sources:law:92,law:93"  # 正常值
    seg.source_id = "sources:law:92,law:WRONG"       # 竄改
    assert claim.split("[")[0] in seg.text or "附款" in seg.text
    assert [s.id for s in seg.sources] == [src_a.id, src_b.id]
    assert seg.traceability == [
        {"kind": "source", "id": src_a.id},
        {"kind": "source", "id": src_b.id},
    ]

    # 閘門必須攔截 source_id 不一致
    with caplog.at_level(logging.WARNING):
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗") as ei:
            require_traceable_note_product(product, source="fail-sourceid-mismatch")
    msg = str(ei.value)
    assert "source_id" in msg, f"應指出 source_id 未對上，實際：{msg!r}"
    assert "note_traceability_failed" in caplog.text
    assert "fail-sourceid-mismatch" in caplog.text
    # 錯誤訊息應含期望值與實際值
    assert "law:93" in msg or "WRONG" in msg


def test_failure_semantic_source_id_mismatch_vs_source_ids_list(caplog):
    """最小負例：source_ids 列表正確，但 source_id 欄位缺少其中一個 ID。

    source_ids = [law:92, law:93]，source_id 應為 sources:law:92,law:93，
    實際却為 sources:law:92 → 閘必須指出 source_id 欄位與 source_ids 不符。
    """
    import logging

    from note_filler.pipeline import require_traceable_note_product

    src_a = _source("law:92", "行政程序法第 92 條")
    src_b = _source("law:93", "行政程序法第 93 條", level="B")
    gap = Gap("附款之限制為何？", "missing", "原稿未列")
    claim = "附款不得違背行政處分之目的[^1][^2]。"
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src_a, src_b]},
        {gap.question: WrittenSupplement(claim, [src_a.id, src_b.id])},
        {gap.question: cross_validate(gap.question, [src_a, src_b])},
    )
    seg = product.segments[-1]
    # 驗證組裝正確後，人為截斷 source_id 欄位
    assert seg.source_id == "sources:law:92,law:93"
    seg.source_id = "sources:law:92"                   # 缺少 law:93
    assert [s.id for s in seg.sources] == [src_a.id, src_b.id]

    with caplog.at_level(logging.WARNING):
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗") as ei:
            require_traceable_note_product(product, source="fail-sourceid-truncated")
    msg = str(ei.value)
    assert "source_id" in msg, f"應指出 source_id 未對上，實際：{msg!r}"
    assert "note_traceability_failed" in caplog.text
    assert "fail-sourceid-truncated" in caplog.text


def test_failure_semantic_source_fragment_empty(caplog):
    """最小負例：論點有文字、來源 ID 與追溯正確，但來源 content 為空白 → 片段缺失。

    必須：
    1) 綁定報告 no_empty_fragments 為 False、binding_status 為 fail
    2) require_traceable_note_product 明確 raise 並指出缺失片段
    """
    import logging

    from note_filler.binding_report import build_binding_report, parse_binding_report
    from note_filler.pipeline import require_traceable_note_product

    gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
    claim = "行政處分應符合法定要件。"
    src_empty = Source(
        id="law:empty-frag",
        title="行政程序法第 92 條",
        url=None, level="A",
        content="",
        fetched_date="2026-07-25",
        doc_date=None, distance=0.5,
    )
    with caplog.at_level(logging.WARNING):
        product = assemble_correction(
            _doc(), [gap],
            {gap.question: [src_empty]},
            {gap.question: WrittenSupplement(claim, [src_empty.id])},
            {gap.question: cross_validate(gap.question, [src_empty])},
        )
    seg = product.segments[-1]
    assert seg.type == "supplement"
    assert claim in seg.text
    assert len(seg.sources) == 1
    assert seg.sources[0].id == "law:empty-frag"
    assert not seg.sources[0].content.strip()

    # 綁定報告必須明確標示 fail（來源片段缺失）
    report = parse_binding_report(build_binding_report(product))
    arg = report["arguments"][0]
    assert arg["binding_status"] == "fail", (
        f"空片段來源之 binding_status 應為 fail，實際：{arg['binding_status']!r}"
    )
    assert arg["binding_ok"] is False
    assert arg["checks"]["no_empty_fragments"] is False
    assert report["summary"]["fail"] == 1
    assert report["summary"]["all_arguments_ok"] is False

    # 閘門必須明確 raise 且錯誤指向缺失片段
    with caplog.at_level(logging.WARNING):
        with pytest.raises(RuntimeError, match="來源追溯驗證失敗") as ei:
            require_traceable_note_product(product, source="fail-empty-frag")
    msg = str(ei.value)
    assert "片段缺失" in msg, (
        f"失敗訊息應指出片段缺失，實際：{msg!r}"
    )
    assert "note_traceability_failed" in caplog.text
    assert "fail-empty-frag" in caplog.text


# ---- 最小驗收：直接讀序列化成品，逐筆論點↔來源 ID／片段／位置 ---------------


def test_serialized_product_each_argument_source_id_maps_1to1_fragment_position(
    tmp_path,
):
    """最小驗收：直接讀取成品序列化輸出，逐筆斷言來源綁定。

    對每個論點（supplement）：
      1. 至少一個來源
      2. 來源 ID 與來源片段（content）／位置（list index）一一對上
      3. sources / traceability / source_id 欄位無重複、無遺漏
    不依賴 LLM／網路；只讀落盤 JSON + binding_report.json。
    """
    import json
    from pathlib import Path

    from note_filler.export import to_json
    from note_filler.pipeline import require_traceable_note_product

    # 每個來源 content 含可機器比對的唯一片段標記，便於 ID↔片段對位
    src_a = Source(
        id="law:92",
        title="行政程序法第 92 條",
        url=None,
        level="A",
        content="【片段:law:92】行政處分，係指行政機關就公法上具體事件所為之決定。",
        fetched_date="2026-07-25",
        doc_date=None,
        distance=0.1,
    )
    src_b1 = Source(
        id="law:93",
        title="行政程序法第 93 條",
        url=None,
        level="A",
        content="【片段:law:93】行政機關作成行政處分有裁量權時，得為附款。",
        fetched_date="2026-07-25",
        doc_date=None,
        distance=0.2,
    )
    src_b2 = Source(
        id="law:94",
        title="行政程序法第 94 條",
        url=None,
        level="B",
        content="【片段:law:94】前條之附款不得違背行政處分之目的。",
        fetched_date="2026-07-25",
        doc_date=None,
        distance=0.3,
    )
    gap_a = Gap("定義？", "missing", "未說明")
    gap_b = Gap("附款限制？", "missing", "未說明")
    product = assemble_correction(
        _doc(),
        [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b1, src_b2]},
        {
            gap_a.question: WrittenSupplement("定義[^1]。", [src_a.id]),
            gap_b.question: WrittenSupplement(
                "限制[^1][^2]。", [src_b1.id, src_b2.id]
            ),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b1, src_b2]),
        },
    )
    require_traceable_note_product(product, source="serialized-1to1-acceptance")

    # 落盤：成品 JSON + 綁定報告（模擬交付後的序列化輸出）
    out_dir = Path(tmp_path)
    product_path = out_dir / "note_product.json"
    product_path.write_text(
        json.dumps(to_json(product), ensure_ascii=False, indent=2),
        encoding="utf-8",
        newline="\n",
    )
    report_path = write_binding_report(out_dir / "out.md", product)
    assert report_path.is_file()
    assert product_path.is_file()

    # ---- 直接讀取序列化輸出（不再使用記憶體中的 product 物件）----
    serialized = json.loads(product_path.read_text(encoding="utf-8"))
    report = parse_binding_report(
        json.loads(report_path.read_text(encoding="utf-8"))
    )

    segments = serialized["segments"]
    arguments = report["arguments"]
    assert len(arguments) >= 1, "至少一個論點"
    assert report["summary"]["all_arguments_ok"] is True
    assert report["summary"]["fail"] == 0

    # 反向索引：每個來源 ID 應能對回至少一個論點 index
    source_usage = report["source_usage"]
    seen_source_ids: list[str] = []

    for arg in arguments:
        seg = segments[arg["segment_index"]]
        assert seg["type"] == "supplement"
        assert arg["argument_text"] == seg["text"]

        source_ids = list(arg["source_ids"])
        trace_ids = list(arg["trace_source_ids"])
        sources = list(seg["sources"])
        trace_refs = [
            t for t in seg["traceability"] if t.get("kind") == "source"
        ]

        # (1) 每個論點至少一個來源
        assert len(source_ids) >= 1, (
            f"論點[{arg['argument_index']}] 缺少來源: {arg['argument_text']!r}"
        )
        assert arg["checks"]["at_least_one_source"] is True
        assert len(sources) >= 1

        # (2) 無重複
        assert len(source_ids) == len(set(source_ids)), (
            f"論點[{arg['argument_index']}] source_ids 重複: {source_ids}"
        )
        assert len(trace_ids) == len(set(trace_ids)), (
            f"論點[{arg['argument_index']}] trace_source_ids 重複: {trace_ids}"
        )
        obj_ids = [s["id"] for s in sources]
        assert len(obj_ids) == len(set(obj_ids)), (
            f"論點[{arg['argument_index']}] sources 物件 id 重複: {obj_ids}"
        )
        assert arg["checks"]["no_duplicate_sources"] is True

        # (3) 無遺漏／無多餘：source_ids ↔ sources ↔ trace 集合與順序一致
        assert obj_ids == source_ids, (
            f"論點[{arg['argument_index']}] sources 物件序與 source_ids 不一致: "
            f"{obj_ids} vs {source_ids}"
        )
        assert trace_ids == source_ids, (
            f"論點[{arg['argument_index']}] trace_source_ids 與 source_ids 不一致: "
            f"{trace_ids} vs {source_ids}"
        )
        assert [t["id"] for t in trace_refs] == source_ids
        assert arg["checks"]["no_omitted_traces"] is True
        assert arg["checks"]["no_extra_traces"] is True
        assert arg["checks"]["source_traceable"] is True

        # source_id 欄位與 ID 列表對齊
        expected_field = f"sources:{','.join(source_ids)}"
        assert seg["source_id"] == expected_field
        assert arg["source_id_field"] == expected_field
        assert arg["checks"]["source_id_field_aligned"] is True

        # (4) 來源 ID ↔ 來源片段／位置 一一對上
        #     位置 = list index；片段 = Source.content 中的唯一標記
        for pos, sid in enumerate(source_ids):
            src_obj = sources[pos]
            assert src_obj["id"] == sid, (
                f"位置 {pos}：期望 id={sid!r}，實際 {src_obj['id']!r}"
            )
            assert trace_refs[pos]["id"] == sid, (
                f"位置 {pos}：trace 未對上 id={sid!r}"
            )
            fragment = src_obj.get("content") or ""
            assert fragment.strip(), (
                f"來源 {sid!r} 缺少可驗證片段 content"
            )
            # 片段標記必須含自身 ID，確保 ID↔片段可對位、不與其他來源混淆
            assert f"【片段:{sid}】" in fragment, (
                f"來源 {sid!r} 的片段未含自身標記；content={fragment!r}"
            )
            # 不得誤掛其他來源的片段標記
            for other in source_ids:
                if other == sid:
                    continue
                assert f"【片段:{other}】" not in fragment, (
                    f"來源 {sid!r} 片段誤含其他 id 標記 {other!r}"
                )

            # 反向索引：此 source_id 的 usage 必須包含本論點
            assert sid in source_usage, f"source_usage 遺漏 {sid!r}"
            assert arg["argument_index"] in source_usage[sid], (
                f"source_usage[{sid!r}] 未含 argument_index="
                f"{arg['argument_index']}"
            )
            seen_source_ids.append(sid)

        assert arg["binding_ok"] is True
        assert arg["binding_status"] == "pass"
        assert arg["source_count"] == len(source_ids)
        assert arg["cardinality"] in ("one_to_one", "one_to_many")

    # 全域：source_usage 的 key 集合 = 所有論點用過的來源（無重複遺漏）
    assert set(source_usage) == set(seen_source_ids)
    # 一對一 + 一對多各至少一筆（混合情境鎖定）
    assert report["summary"]["one_to_one"] >= 1
    assert report["summary"]["one_to_many"] >= 1


def test_product_output_each_argument_has_binding_and_dual_necessity_views(
    tmp_path,
):
    """成品逐筆保留來源綁定與功能缺口／使用者價值，並可解析 1:1、1:N。"""
    import json

    from note_filler.export import to_json

    src_a = _source("law:92", "行政程序法第 92 條")
    src_b = _source("law:93", "行政程序法第 93 條")
    src_c = _source("law:94", "行政程序法第 94 條", level="B")
    gap_a = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    gap_b = Gap("附款有何限制？", "missing", "原稿未說明附款限制")
    product = assemble_correction(
        _doc(),
        [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b, src_c]},
        {
            gap_a.question: WrittenSupplement("行政處分定義[^1]。", [src_a.id]),
            gap_b.question: WrittenSupplement(
                "附款限制[^1][^2]。", [src_b.id, src_c.id]
            ),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b, src_c]),
        },
    )
    # assemble 已寫入 functional_gap／user_value；此處覆寫為更具體的使用者價值敘述，
    # 驗證序列化後雙視角仍完整保留（非空且與成品一致）。
    for segment, user_value in zip(
        product.segments[1:],
        ("讓讀者辨識行政處分的適用範圍", "讓讀者判斷附款是否合法"),
        strict=True,
    ):
        segment.user_value = user_value
        assert segment.functional_gap.strip()

    product_path = tmp_path / "note_product.json"
    product_path.write_text(
        json.dumps(to_json(product), ensure_ascii=False, indent=2),
        encoding="utf-8",
        newline="\n",
    )
    report_path = write_binding_report(tmp_path / "note_product.md", product)

    serialized = json.loads(product_path.read_text(encoding="utf-8"))
    report = parse_binding_report(
        json.loads(report_path.read_text(encoding="utf-8"))
    )
    expected = [
        ("one_to_one", ["law:92"]),
        ("one_to_many", ["law:93", "law:94"]),
    ]

    assert report["argument_count"] == len(expected)
    assert report["summary"]["one_to_one"] == 1
    assert report["summary"]["one_to_many"] == 1
    for argument, (cardinality, source_ids) in zip(
        report["arguments"], expected, strict=True
    ):
        segment = serialized["segments"][argument["segment_index"]]
        assert segment["argument_id"] == f"argument:{argument['argument_index']}"
        assert argument["argument_text"] == segment["text"]
        assert argument["cardinality"] == segment["cardinality"] == cardinality
        assert argument["source_ids"] == segment["source_ids"] == source_ids
        assert argument["binding_ok"] is True
        # at_least_two_openable_links 為附加輸出規則，不影響 binding_ok
        core_checks = {
            k: v for k, v in argument["checks"].items()
            if k != "at_least_two_openable_links"
        }
        assert all(core_checks.values())
        for field in ("functional_gap", "user_value"):
            assert argument[field] == segment[field]
            assert argument[field].strip(), (
                f"論點[{argument['argument_index']}] 缺少 {field}"
            )

    assert report["source_usage"] == {
        "law:92": [0],
        "law:93": [1],
        "law:94": [1],
    }


def test_positive_disk_product_multi_angle_and_one_to_many_four_way_consistency(
    tmp_path,
):
    """正向：直接讀成品／綁定報告，至少一則論點同時滿足多有效角度與一對多來源。

    驗收條件：
      1. 落盤後只靠 JSON 讀回（不依賴記憶體 product）
      2. 筆記層有 ≥2 個不同且不重複的有效角度（definition + limitation）
      3. 至少一則論點 cardinality=one_to_many 且 source_count≥2
      4. 工具鏈可逐筆核對：角度、來源、功能缺口、使用者價值 四者
         在成品 segments 與 binding_report.arguments 完全一致
    """
    import json
    from pathlib import Path

    from note_filler.export import to_json
    from note_filler.pipeline import require_traceable_note_product

    src_def = Source(
        id="law:92",
        title="行政程序法第 92 條",
        url=None,
        level="A",
        content="【片段:law:92】行政處分，係指行政機關就公法上具體事件所為之決定。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.1,
    )
    src_lim_a = Source(
        id="law:93",
        title="行政程序法第 93 條",
        url=None,
        level="A",
        content="【片段:law:93】行政機關作成行政處分有裁量權時，得為附款。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.2,
    )
    src_lim_b = Source(
        id="law:94",
        title="行政程序法第 94 條",
        url=None,
        level="B",
        content="【片段:law:94】前條之附款不得違背行政處分之目的。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.3,
    )
    # 同一主題「行政處分」：定義角度（1:1）＋限制角度且一對多來源（1:N）
    gap_def = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    gap_lim = Gap("行政處分有何限制？", "missing", "原稿未說明附款限制")
    product = assemble_correction(
        _doc(),
        [gap_def, gap_lim],
        {
            gap_def.question: [src_def],
            gap_lim.question: [src_lim_a, src_lim_b],
        },
        {
            gap_def.question: WrittenSupplement(
                "行政處分定義參照[^1]。", [src_def.id]
            ),
            gap_lim.question: WrittenSupplement(
                "附款限制須兼顧目的[^1][^2]。",
                [src_lim_a.id, src_lim_b.id],
            ),
        },
        {
            gap_def.question: cross_validate(gap_def.question, [src_def]),
            gap_lim.question: cross_validate(
                gap_lim.question, [src_lim_a, src_lim_b]
            ),
        },
    )
    # 鎖定明確的功能缺口／使用者價值，便於逐筆四向比對
    expected_necessity = [
        (
            "原稿未定義行政處分",
            "補齊讀者對「行政處分如何定義？」所需的說明",
        ),
        (
            "原稿未說明附款限制",
            "補齊讀者對「行政處分有何限制？」所需的說明",
        ),
    ]
    for segment, (fg, uv) in zip(product.segments[1:], expected_necessity, strict=True):
        segment.functional_gap = fg
        segment.user_value = uv

    require_traceable_note_product(
        product, source="positive-multi-angle-one-to-many"
    )

    out_dir = Path(tmp_path)
    product_path = out_dir / "note_product.json"
    product_path.write_text(
        json.dumps(to_json(product), ensure_ascii=False, indent=2),
        encoding="utf-8",
        newline="\n",
    )
    report_path = write_binding_report(out_dir / "out.md", product)
    assert product_path.is_file()
    assert report_path.is_file()

    # ---- 只讀落盤：成品 JSON + binding_report.json ----
    serialized = json.loads(product_path.read_text(encoding="utf-8"))
    report = parse_binding_report(
        json.loads(report_path.read_text(encoding="utf-8"))
    )

    assert report["summary"]["all_arguments_ok"] is True
    assert report["summary"]["fail"] == 0
    assert report["argument_count"] == 2

    # 多個不同且不重複的有效角度（筆記層）
    angle_summary = report["angle_coverage_summary"]
    assert set(angle_summary["unique_angle_types"]) == {"definition", "limitation"}
    assert angle_summary["effective_angle_count"] == 2
    assert angle_summary["required_effective_angle_count"] == 2
    assert angle_summary["excluded_angle_count"] == 0
    assert angle_summary["duplicate_pairs"] == []
    assert angle_summary["synonym_pairs"] == []
    assert angle_summary["has_sufficient_angles"] is True
    assert angle_summary["coverage_ok"] is True
    assert serialized["angle_coverage_summary"]["coverage_ok"] is True
    assert serialized["angle_coverage_summary"]["effective_angle_count"] == 2

    # 至少一則論點為來源一對多
    one_to_many_args = [
        a for a in report["arguments"] if a["cardinality"] == "one_to_many"
    ]
    assert len(one_to_many_args) >= 1, "須至少一則論點為 one_to_many"
    assert report["summary"]["one_to_many"] >= 1
    for otm in one_to_many_args:
        assert otm["source_count"] >= 2
        assert len(otm["source_ids"]) >= 2
        assert len(otm["source_ids"]) == len(set(otm["source_ids"]))
        assert otm["binding_ok"] is True
        assert otm["binding_status"] == "pass"

    # 該一對多論點同時位於多有效角度集合中（限制角度 + 多源）
    otm = one_to_many_args[0]
    assert otm["angle_coverage"]["angle_type"] == "limitation"
    assert otm["angle_coverage"]["relation"]["kind"] == "unique"
    assert otm["angle_coverage"]["effective_angle_count"] == 1
    assert otm["source_ids"] == ["law:93", "law:94"]

    # 逐筆核對：角度、來源、功能缺口、使用者價值 四者一致
    angle_types_seen: set[str] = set()
    angle_keys_seen: set[str] = set()
    for arg in report["arguments"]:
        seg = serialized["segments"][arg["segment_index"]]
        assert seg["type"] == "supplement"
        ac_report = arg["angle_coverage"]
        ac_product = seg["angle_coverage"]

        # (1) 角度：成品 ↔ 報告
        assert ac_report["angle_type"] == ac_product["angle_type"] == seg["angle_type"]
        assert ac_report["angle_key"] == ac_product["angle_key"] == seg["angle_key"]
        assert list(ac_report["angle_labels"]) == list(ac_product["angle_labels"])
        assert list(ac_report["angle_labels"]) == list(seg["angle_labels"])
        assert list(ac_report["covered_facets"]) == list(ac_product["covered_facets"])
        assert "angle:" + ac_report["angle_type"] in ac_report["covered_facets"]
        assert "necessity:functional_gap" in ac_report["covered_facets"]
        assert "necessity:user_value" in ac_report["covered_facets"]
        assert "question" in ac_report["covered_facets"]
        assert ac_report["relation"]["kind"] == "unique"
        assert ac_report["effective_angle_count"] == 1
        assert ac_report["duplicate_exclusion"]["excluded"] is False
        assert arg["checks"]["has_angle_coverage"] is True
        assert arg["checks"]["meets_angle_coverage_threshold"] is True
        assert arg["checks"]["angle_facet_complete"] is True
        angle_types_seen.add(ac_report["angle_type"])
        angle_keys_seen.add(ac_report["angle_key"])

        # (2) 來源：成品 ↔ 報告（含片段／位置）
        source_ids = list(arg["source_ids"])
        assert source_ids == list(seg["source_ids"])
        assert source_ids == list(arg["trace_source_ids"])
        assert arg["cardinality"] == seg["cardinality"]
        assert arg["source_count"] == len(source_ids) >= 1
        sources = list(seg["sources"])
        assert [s["id"] for s in sources] == source_ids
        for pos, sid in enumerate(source_ids):
            assert sources[pos]["id"] == sid
            fragment = sources[pos].get("content") or ""
            assert fragment.strip(), f"來源 {sid!r} 片段空白"
            assert f"【片段:{sid}】" in fragment
        assert arg["checks"]["at_least_one_source"] is True
        assert arg["checks"]["source_traceable"] is True
        assert arg["checks"]["no_duplicate_sources"] is True
        assert arg["checks"]["no_omitted_traces"] is True
        assert arg["checks"]["no_extra_traces"] is True
        assert arg["checks"]["no_empty_fragments"] is True

        # (3) 功能缺口 + (4) 使用者價值：成品 ↔ 報告
        assert arg["functional_gap"] == seg["functional_gap"]
        assert arg["user_value"] == seg["user_value"]
        assert arg["functional_gap"].strip()
        assert arg["user_value"].strip()
        assert arg["checks"]["has_functional_gap"] is True
        assert arg["checks"]["has_user_value"] is True
        assert arg["checks"]["angle_functional_gap_present"] is True
        assert arg["checks"]["angle_user_value_present"] is True

        # 四者齊備且綁定通過
        assert arg["binding_ok"] is True
        assert arg["binding_status"] == "pass"
        # at_least_two_openable_links 為附加輸出規則，不影響 binding_ok
        core_checks = {
            k: v for k, v in arg["checks"].items()
            if k != "at_least_two_openable_links"
        }
        assert all(core_checks.values())

    # 跨論點：有效角度真的不同且不重複
    assert angle_types_seen == {"definition", "limitation"}
    assert len(angle_keys_seen) == 2

    # 必要性具體值與寫入時一致
    for arg, (fg, uv) in zip(report["arguments"], expected_necessity, strict=True):
        assert arg["functional_gap"] == fg
        assert arg["user_value"] == uv

    # 反向索引：一對多來源皆指向限制論點
    assert report["source_usage"] == {
        "law:92": [0],
        "law:93": [1],
        "law:94": [1],
    }


# ---- 三元驗收：功能缺口 + 使用者價值 + 關聯知識 同論點同來源呼應 -----------


_RE_VISIBLE_SUMMARY = re.compile(
    r"argument_id=(?P<argument_id>[^；]+)"
    r"；functional_gap=(?P<functional_gap>.*?)"
    r"；user_value=(?P<user_value>.*?)"
    r"；related_knowledge=(?P<related_knowledge>.*)$"
)
_RE_SOURCE_LIST = re.compile(
    # Markdown 可能為 **來源清單**：ids（card）— 冒號前允許粗體星號
    r"來源清單\*{0,2}[：:]\s*(?P<ids>[^（(]+)"
    r"(?:[（(](?P<cardinality>[^）)]+)[）)])?"
)


def _parse_md_argument_blocks(md: str) -> list[dict[str, str]]:
    """從人類可讀 Markdown 逐段解析補充論點的三元與來源。"""
    blocks: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for raw in md.splitlines():
        line = raw.strip()
        if line.startswith("> 【補充】") or line.startswith("【補充】"):
            if current is not None:
                blocks.append(current)
            current = {
                "summary": "",
                "functional_gap": "",
                "user_value": "",
                "related_knowledge": "",
                "source_ids": "",
                "cardinality": "",
                "argument_id": "",
                "visible_summary": "",
            }
            # 去掉前綴與 footnote 標記，留下摘要正文
            body = re.sub(r"^>\s*", "", line)
            body = re.sub(r"^【補充】(?:⚠待補證\s*)?", "", body)
            body = re.sub(r"\[\^\d+\]", "", body).strip()
            current["summary"] = body
            continue
        if current is None:
            continue
        plain = re.sub(r"^>\s*", "", line)
        if plain.startswith("**功能缺口**：") or plain.startswith("功能缺口："):
            current["functional_gap"] = plain.split("：", 1)[1].strip()
        elif plain.startswith("**使用者價值**：") or plain.startswith("使用者價值："):
            current["user_value"] = plain.split("：", 1)[1].strip()
        elif plain.startswith("**關聯知識**：") or plain.startswith("關聯知識："):
            current["related_knowledge"] = plain.split("：", 1)[1].strip()
        elif plain.startswith("**來源清單**：") or plain.startswith("來源清單"):
            m = _RE_SOURCE_LIST.search(plain)
            if m:
                current["source_ids"] = m.group("ids").strip()
                current["cardinality"] = (m.group("cardinality") or "").strip()
        elif plain.startswith("**摘要可見**：") or plain.startswith("摘要可見："):
            visible = plain.split("：", 1)[1].strip()
            current["visible_summary"] = visible
            vm = _RE_VISIBLE_SUMMARY.fullmatch(visible)
            assert vm is not None, f"摘要可見格式不可解析: {visible!r}"
            # 以摘要可見內嵌三元覆寫，確保同列、同 argument_id 一致
            current["argument_id"] = vm.group("argument_id").strip()
            current["functional_gap"] = vm.group("functional_gap").strip()
            current["user_value"] = vm.group("user_value").strip()
            current["related_knowledge"] = vm.group("related_knowledge").strip()
        elif plain.startswith("**論點ID**：") or plain.startswith("論點ID："):
            current["argument_id"] = plain.split("：", 1)[1].strip()
    if current is not None:
        blocks.append(current)
    return blocks


def _shared_topic_token(summary: str, functional_gap: str, user_value: str) -> str | None:
    """找出三元文字共同出現的關鍵詞（至少 2 字），作為互相呼應的最小證據。"""
    runs = re.findall(r"[\u4e00-\u9fff]+", f"{functional_gap} {user_value}")
    candidates: list[str] = []
    seen: set[str] = set()
    for run in runs:
        # 由長到短產生 2..min(8,len) 的子字串，避免整段貪婪匹配吞掉主題詞
        for n in range(min(8, len(run)), 1, -1):
            for i in range(0, len(run) - n + 1):
                token = run[i : i + n]
                if token not in seen:
                    seen.add(token)
                    candidates.append(token)
    candidates.sort(key=len, reverse=True)
    for token in candidates:
        if token in summary and token in functional_gap and token in user_value:
            return token
    return None


def _assert_triad_acceptance(
    argument_id: str,
    summary: str,
    functional_gap: str,
    user_value: str,
    related_knowledge: str = "",
) -> None:
    """驗收三元必填且摘要／關聯知識確實呼應功能缺口與使用者價值。

    related_knowledge 若提供，必須明示「支撐決策品質」與「補強使用者理解」。
    """
    fields = {
        "summary": summary,
        "functional_gap": functional_gap,
        "user_value": user_value,
    }
    missing = [name for name, value in fields.items() if not value.strip()]
    if missing:
        raise AssertionError(f"{argument_id}：缺少欄位 {'、'.join(missing)}")
    if _shared_topic_token(summary, functional_gap, user_value) is None:
        raise AssertionError(
            f"{argument_id}：內容不一致位置 summary 未反映 functional_gap／user_value；"
            f"summary={summary!r}；functional_gap={functional_gap!r}；user_value={user_value!r}"
        )
    if related_knowledge.strip():
        if "支撐決策品質" not in related_knowledge:
            raise AssertionError(
                f"{argument_id}：related_knowledge 未說明如何支撐決策品質；"
                f"related_knowledge={related_knowledge!r}"
            )
        if "補強使用者理解" not in related_knowledge:
            raise AssertionError(
                f"{argument_id}：related_knowledge 未說明如何補強使用者理解；"
                f"related_knowledge={related_knowledge!r}"
            )
        if _shared_topic_token(related_knowledge, functional_gap, user_value) is None:
            raise AssertionError(
                f"{argument_id}：related_knowledge 未與 functional_gap／user_value 呼應；"
                f"related_knowledge={related_knowledge!r}"
            )


def test_final_output_each_argument_triad_coheres_with_same_source_and_argument(
    tmp_path,
):
    """最終輸出：每個論點同時含摘要、功能缺口、使用者價值，三者呼應且同來源。

    只讀落盤成品（note_product.json / binding_report.json / .md），
    逐筆對應同一 argument_id 與 source_ids；涵蓋 1:1 與 1:N。
    """
    import json
    from pathlib import Path

    from note_filler.export import to_json, to_markdown
    from note_filler.pipeline import require_traceable_note_product

    src_a = Source(
        id="law:92",
        title="行政程序法第 92 條",
        url=None,
        level="A",
        content="【片段:law:92】行政處分，係指行政機關就公法上具體事件所為之決定。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.1,
    )
    src_b1 = Source(
        id="law:93",
        title="行政程序法第 93 條",
        url=None,
        level="A",
        content="【片段:law:93】行政機關作成行政處分有裁量權時，得為附款。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.2,
    )
    src_b2 = Source(
        id="law:94",
        title="行政程序法第 94 條",
        url=None,
        level="B",
        content="【片段:law:94】前條之附款不得違背行政處分之目的。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.3,
    )
    gap_a = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    gap_b = Gap("附款有何限制？", "missing", "原稿未說明附款限制")
    claim_a = "行政處分定義參照[^1]。"
    claim_b = "附款限制須兼顧目的[^1][^2]。"
    product = assemble_correction(
        _doc(),
        [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b1, src_b2]},
        {
            gap_a.question: WrittenSupplement(claim_a, [src_a.id]),
            gap_b.question: WrittenSupplement(claim_b, [src_b1.id, src_b2.id]),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b1, src_b2]),
        },
    )
    # 鎖定三元期望值（摘要去掉 footnote 標記後與 argument_text 對齊）
    expected = [
        {
            "argument_id": "argument:0",
            "summary": "行政處分定義參照。",
            "functional_gap": "原稿未定義行政處分",
            "user_value": "補齊讀者對「行政處分如何定義？」所需的說明",
            "source_ids": ["law:92"],
            "cardinality": "one_to_one",
            "topic": "行政處分",
        },
        {
            "argument_id": "argument:1",
            "summary": "附款限制須兼顧目的。",
            "functional_gap": "原稿未說明附款限制",
            "user_value": "補齊讀者對「附款有何限制？」所需的說明",
            "source_ids": ["law:93", "law:94"],
            "cardinality": "one_to_many",
            "topic": "附款",
        },
    ]
    for segment, exp in zip(product.segments[1:], expected, strict=True):
        segment.functional_gap = exp["functional_gap"]
        segment.user_value = exp["user_value"]
        assert segment.argument_id == exp["argument_id"]

    require_traceable_note_product(product, source="triad-coherence-acceptance")

    out_dir = Path(tmp_path)
    product_path = out_dir / "note_product.json"
    md_path = out_dir / "note_product.md"
    product_path.write_text(
        json.dumps(to_json(product), ensure_ascii=False, indent=2),
        encoding="utf-8",
        newline="\n",
    )
    md_path.write_text(to_markdown(product), encoding="utf-8", newline="\n")
    report_path = write_binding_report(md_path, product)
    assert product_path.is_file()
    assert md_path.is_file()
    assert report_path.is_file()

    # ---- 只讀落盤 ----
    serialized = json.loads(product_path.read_text(encoding="utf-8"))
    report = parse_binding_report(
        json.loads(report_path.read_text(encoding="utf-8"))
    )
    md = md_path.read_text(encoding="utf-8")
    md_blocks = _parse_md_argument_blocks(md)

    assert report["summary"]["all_arguments_ok"] is True
    assert report["summary"]["fail"] == 0
    assert report["argument_count"] == 2
    assert report["summary"]["one_to_one"] == 1
    assert report["summary"]["one_to_many"] == 1
    assert len(md_blocks) == 2

    for arg, exp, block in zip(report["arguments"], expected, md_blocks, strict=True):
        seg = serialized["segments"][arg["segment_index"]]
        assert seg["type"] == "supplement"
        assert arg["argument_id"] == exp["argument_id"] == seg["argument_id"]
        assert block["argument_id"] == exp["argument_id"]

        # (1) 三元齊備：摘要 / 功能缺口 / 使用者價值 / 關聯知識
        summary = (arg["argument_text"] or "").strip()
        # 摘要以正文為準（可能含 footnote 標記）；比對時去掉 [^n]
        summary_plain = re.sub(r"\[\^\d+\]", "", summary).strip()
        fg = (arg["functional_gap"] or "").strip()
        uv = (arg["user_value"] or "").strip()
        rk = (arg["related_knowledge"] or "").strip()
        _assert_triad_acceptance(
            exp["argument_id"], summary_plain, fg, uv, related_knowledge=rk
        )
        assert "（未提供）" not in fg and "（未提供）" not in uv
        assert "支撐決策品質" in rk and "補強使用者理解" in rk

        # 成品 segment ↔ binding_report 同論點一致
        assert re.sub(r"\[\^\d+\]", "", seg["text"]).strip() == summary_plain
        assert seg["functional_gap"] == fg == exp["functional_gap"]
        assert seg["user_value"] == uv == exp["user_value"]
        assert seg["related_knowledge"] == rk
        assert arg["checks"]["has_functional_gap"] is True
        assert arg["checks"]["has_user_value"] is True
        assert arg["checks"]["has_related_knowledge"] is True
        assert arg["checks"]["related_knowledge_consistent"] is True

        # (2) 互相呼應：主題詞同時出現在三元中
        assert exp["topic"] in summary_plain
        assert exp["topic"] in fg
        assert exp["topic"] in uv
        assert exp["topic"] in rk
        # user_value 須回扣同一問題（與摘要／缺口同論點）
        assert "補齊讀者對「" in uv and "」所需的說明" in uv

        # (3) 同一來源：product / report / markdown 來源清單對齊
        source_ids = list(arg["source_ids"])
        assert source_ids == exp["source_ids"]
        assert source_ids == list(seg["source_ids"])
        assert source_ids == list(arg["trace_source_ids"])
        assert [s["id"] for s in seg["sources"]] == source_ids
        assert arg["cardinality"] == exp["cardinality"] == seg["cardinality"]
        assert arg["source_count"] == len(source_ids)
        assert arg["binding_ok"] is True
        assert arg["binding_status"] == "pass"
        assert arg["checks"]["at_least_one_source"] is True
        assert arg["checks"]["source_traceable"] is True

        # Markdown 摘要可見列內嵌三元，且與獨立欄位同值同論點
        block_summary_plain = re.sub(r"\[\^\d+\]", "", block["summary"]).strip()
        assert block_summary_plain == summary_plain
        assert block_summary_plain == exp["summary"]
        assert block["functional_gap"] == fg
        assert block["user_value"] == uv
        assert block["related_knowledge"] == rk
        assert block["source_ids"].replace(" ", "") == ",".join(source_ids)
        assert "（未提供）" not in block["visible_summary"]
        assert f"argument_id={exp['argument_id']}" in block["visible_summary"]
        # related_knowledge 與同一 argument_id 綁定，明示決策品質／使用者理解
        vis_summary_m = re.search(
            r"；related_knowledge=(.*)$", block["visible_summary"]
        )
        assert vis_summary_m is not None
        assert vis_summary_m.group(1).strip() == rk
        assert "支撐決策品質" in vis_summary_m.group(1)
        assert "補強使用者理解" in vis_summary_m.group(1)
        assert f"functional_gap={fg}" in block["visible_summary"]
        assert f"user_value={uv}" in block["visible_summary"]

    # 反向：來源用法無重複遺漏
    assert report["source_usage"] == {
        "law:92": [0],
        "law:93": [1],
        "law:94": [1],
    }


@pytest.mark.parametrize(
    "missing_field",
    ["functional_gap", "user_value", "summary", "related_knowledge"],
)
def test_final_output_triad_missing_any_field_fails_explicitly(
    missing_field: str,
    tmp_path,
):
    """負例：來源綁定存在，但摘要／功能缺口／使用者價值／關聯知識缺一即明確失敗。

    - functional_gap / user_value / related_knowledge 空欄或缺決策說明
      → binding_ok=fail，checks 點名缺失；write 不得默默通過
    - 摘要（argument_text）空白 → parse/write 拒絕空欄 argument_text，
      不得默默產出可通過的最終報告
    """
    from pathlib import Path

    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("行政處分定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    assert seg.type == "supplement"
    assert [s.id for s in seg.sources] == ["law:92"]

    if missing_field == "functional_gap":
        seg.functional_gap = ""
    elif missing_field == "user_value":
        seg.user_value = "   "
    elif missing_field == "related_knowledge":
        # 故意寫入無決策品質／使用者理解說明的殘缺文字
        seg.related_knowledge = "僅有摘要卻未說明決策與理解價值"
    else:
        seg.text = "   "

    raw_report = build_binding_report(product)
    arg = raw_report["arguments"][0]
    # 來源層仍可追溯（避免只驗來源而漏掉必要性／摘要）
    assert arg["checks"]["at_least_one_source"] is True
    assert arg["source_ids"] == ["law:92"]

    out = Path(tmp_path) / "out.md"

    if missing_field == "functional_gap":
        assert arg["checks"]["has_functional_gap"] is False
        assert arg["binding_ok"] is False
        assert arg["binding_status"] == "fail"
        assert raw_report["summary"]["all_arguments_ok"] is False
        assert not (arg["functional_gap"] or "").strip()
        assert (arg["argument_text"] or "").strip()
        assert (arg["user_value"] or "").strip()
        # 仍可 parse（空 functional_gap 以 checks 標 fail，非 schema 拒收）
        parsed = parse_binding_report(raw_report)
        assert parsed["arguments"][0]["checks"]["has_functional_gap"] is False
        with pytest.raises(RuntimeError, match="角度有效性驗收失敗|functional_gap") as ei:
            write_binding_report(out, product)
        assert "functional_gap" in str(ei.value)
    elif missing_field == "user_value":
        assert arg["checks"]["has_user_value"] is False
        assert arg["binding_ok"] is False
        assert arg["binding_status"] == "fail"
        assert raw_report["summary"]["all_arguments_ok"] is False
        assert not (arg["user_value"] or "").strip()
        assert (arg["argument_text"] or "").strip()
        assert (arg["functional_gap"] or "").strip()
        parsed = parse_binding_report(raw_report)
        assert parsed["arguments"][0]["checks"]["has_user_value"] is False
        with pytest.raises(RuntimeError, match="角度有效性驗收失敗|user_value") as ei:
            write_binding_report(out, product)
        assert "user_value" in str(ei.value)
    elif missing_field == "related_knowledge":
        assert arg["checks"]["has_related_knowledge"] is False
        assert arg["binding_ok"] is False
        assert arg["binding_status"] == "fail"
        assert raw_report["summary"]["all_arguments_ok"] is False
        assert "支撐決策品質" not in (arg["related_knowledge"] or "")
        assert (arg["argument_text"] or "").strip()
        assert (arg["functional_gap"] or "").strip()
        assert (arg["user_value"] or "").strip()
        parsed = parse_binding_report(raw_report)
        assert parsed["arguments"][0]["checks"]["has_related_knowledge"] is False
        with pytest.raises(RuntimeError, match="關聯知識驗收失敗|related_knowledge") as ei:
            write_binding_report(out, product)
        assert "related_knowledge" in str(ei.value)
    else:
        # 摘要空白：schema 層拒絕 argument_text 空欄
        assert not (arg["argument_text"] or "").strip()
        assert (arg["functional_gap"] or "").strip()
        assert (arg["user_value"] or "").strip()
        with pytest.raises(ValueError, match="argument_text 不可為空欄"):
            parse_binding_report(raw_report)
        with pytest.raises(ValueError, match="argument_text 不可為空欄"):
            write_binding_report(out, product)


def test_triad_acceptance_rejects_summary_with_missing_necessity_fields():
    """負例：摘要存在，但功能缺口與使用者價值皆缺失。"""
    with pytest.raises(AssertionError) as exc_info:
        _assert_triad_acceptance("argument:0", "行政處分定義。", "", "   ")

    assert str(exc_info.value) == (
        "argument:0：缺少欄位 functional_gap、user_value"
    )


def test_triad_acceptance_rejects_summary_not_reflecting_both_perspectives():
    """負例：兩個必要性視角存在，但摘要內容與兩者不一致。"""
    with pytest.raises(AssertionError) as exc_info:
        _assert_triad_acceptance(
            "argument:0",
            "咖啡豆保存方式。",
            "原稿未定義行政處分",
            "補齊讀者對行政處分定義及適用範圍的理解",
        )

    assert (
        "argument:0：內容不一致位置 summary 未反映 functional_gap／user_value"
        in str(exc_info.value)
    )


def test_missing_related_knowledge_with_valid_gap_and_user_value_fails_explicitly(tmp_path):
    """負例：有功能缺口與使用者價值但缺少關聯知識 → 明確失敗並指出缺失欄位。

    最小案例：來源綁定正常、功能缺口與使用者價值皆存在，
    但 related_knowledge 為空或無決策品質／使用者理解說明。
    驗證：
    1) binding_ok=False 且 checks.has_related_knowledge=False
    2) 錯誤訊息明確指出 related_knowledge 缺失
    3) write_binding_report 拒絕輸出
    """
    from pathlib import Path

    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("行政處分定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    assert seg.type == "supplement"
    assert [s.id for s in seg.sources] == ["law:92"]

    # 設定功能缺口與使用者價值，但故意清空關聯知識
    seg.functional_gap = "原稿未定義行政處分，導致讀者無法理解其法律要件"
    seg.user_value = "補齊讀者對行政處分定義及適用範圍的理解"
    seg.related_knowledge = ""  # 缺少關聯知識

    raw_report = build_binding_report(product)
    arg = raw_report["arguments"][0]

    # 來源層仍可追溯
    assert arg["checks"]["at_least_one_source"] is True
    assert arg["source_ids"] == ["law:92"]

    # 功能缺口與使用者價值存在
    assert arg["checks"]["has_functional_gap"] is True
    assert arg["checks"]["has_user_value"] is True
    assert arg["functional_gap"].strip()
    assert arg["user_value"].strip()

    # 關聯知識缺失 → 必須明確失敗
    assert arg["checks"]["has_related_knowledge"] is False
    assert arg["binding_ok"] is False
    assert arg["binding_status"] == "fail"
    assert raw_report["summary"]["all_arguments_ok"] is False

    # parse 階段應能識別缺失
    parsed = parse_binding_report(raw_report)
    assert parsed["arguments"][0]["checks"]["has_related_knowledge"] is False

    # write 階段應拒絕並指出缺失
    out = Path(tmp_path) / "out.md"
    with pytest.raises(RuntimeError, match="關聯知識驗收失敗|related_knowledge") as ei:
        write_binding_report(out, product)
    assert "related_knowledge" in str(ei.value)


def test_related_knowledge_missing_impact_or_importance_fails_explicitly(tmp_path):
    """負例：有關聯知識但未說明影響誰或為何重要 → 明確失敗並指出不一致位置。

    最小案例：關聯知識欄位存在，但內容未說明「影響誰」或「為何重要」，
    即缺少「支撐決策品質」或「補強使用者理解」關鍵詞。
    驗證：
    1) binding_ok=False 且 checks.has_related_knowledge=False
    2) 錯誤訊息指出關聯知識未說明決策品質或使用者理解
    3) write_binding_report 拒絕輸出
    """
    from pathlib import Path

    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("行政處分定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    assert seg.type == "supplement"
    assert [s.id for s in seg.sources] == ["law:92"]

    # 設定功能缺口與使用者價值
    seg.functional_gap = "原稿未定義行政處分，導致讀者無法理解其法律要件"
    seg.user_value = "補齊讀者對行政處分定義及適用範圍的理解"

    # 關聯知識存在但未說明影響誰或為何重要（缺少關鍵詞）
    seg.related_knowledge = "行政處分是一個法律概念，有相關法條規範"

    raw_report = build_binding_report(product)
    arg = raw_report["arguments"][0]

    # 來源層仍可追溯
    assert arg["checks"]["at_least_one_source"] is True
    assert arg["source_ids"] == ["law:92"]

    # 功能缺口與使用者價值存在
    assert arg["checks"]["has_functional_gap"] is True
    assert arg["checks"]["has_user_value"] is True
    assert arg["functional_gap"].strip()
    assert arg["user_value"].strip()

    # 關聯知識存在但未說明決策品質或使用者理解 → 必須明確失敗
    assert arg["checks"]["has_related_knowledge"] is False
    assert arg["binding_ok"] is False
    assert arg["binding_status"] == "fail"
    assert raw_report["summary"]["all_arguments_ok"] is False
    assert "支撐決策品質" not in (arg["related_knowledge"] or "")
    assert "補強使用者理解" not in (arg["related_knowledge"] or "")

    # parse 階段應能識別缺失
    parsed = parse_binding_report(raw_report)
    assert parsed["arguments"][0]["checks"]["has_related_knowledge"] is False

    # write 階段應拒絕並指出缺失
    out = Path(tmp_path) / "out.md"
    with pytest.raises(RuntimeError, match="關聯知識驗收失敗|related_knowledge") as ei:
        write_binding_report(out, product)
    assert "related_knowledge" in str(ei.value)


def test_unrelated_triad_fails_cross_field_consistency_gate(tmp_path):
    """負例：三欄皆存在但主題各寫各的，不得形成可交付的多視角論點。"""
    from pathlib import Path

    src = _source("law:92", "行政程序法第 92 條")
    gap = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: [src]},
        {gap.question: WrittenSupplement("行政處分定義[^1]。", [src.id])},
        {gap.question: cross_validate(gap.question, [src])},
    )
    seg = product.segments[1]
    seg.user_value = "協助咖啡愛好者掌握咖啡豆保存期限"
    seg.related_knowledge = (
        "颱風警報應及早發布（如何支撐決策品質：依風速規劃撤離路線；"
        "如何補強使用者理解：協助居民辨識避難時機）"
    )

    raw_report = build_binding_report(product)
    arg = raw_report["arguments"][0]
    assert arg["functional_gap"].strip()
    assert arg["user_value"].strip()
    assert arg["related_knowledge"].strip()
    assert arg["checks"]["has_related_knowledge"] is True
    assert arg["checks"]["related_knowledge_consistent"] is False
    assert arg["checks"]["at_least_one_source"] is True
    assert arg["binding_ok"] is False
    assert arg["binding_status"] == "fail"
    assert raw_report["summary"]["all_arguments_ok"] is False

    parsed = parse_binding_report(raw_report)
    assert parsed["arguments"][0]["checks"]["related_knowledge_consistent"] is False

    with pytest.raises(RuntimeError, match="跨欄位一致性驗收失敗") as exc_info:
        write_binding_report(Path(tmp_path) / "out.md", product)
    assert "related_knowledge 未對應 functional_gap／user_value" in str(exc_info.value)

    # 報告不可把量測結果竄改為 True 後混過嚴格解析器。
    raw_report["arguments"][0]["checks"]["related_knowledge_consistent"] = True
    with pytest.raises(ValueError, match="跨欄位一致性量測不一致"):
        parse_binding_report(raw_report)


# ---- 四要素驗收：缺口、影響對象、重要性、關聯知識同論點呼應 -------------

def test_final_output_each_argument_has_four_coherent_elements(tmp_path):
    """正向驗收測試：每個論點同時包含缺口、影響對象、重要性、關聯知識，且內容互相呼應。
    
    驗證最終成品中每個論點都具備：
    1. 缺口（functional_gap）：描述原稿缺失的功能性缺口
    2. 影響對象（user_value 中體現）：描述受影響的使用者群體或對象
    3. 重要性（functional_gap/user_value 中體現）：描述補齊後的重要性或價值
    4. 關聯知識（related_knowledge）：描述如何支撐決策品質與補強使用者理解
    
    且四者內容能互相呼應，可逐筆對應到同一 argument_id。
    """
    import json
    from pathlib import Path

    from note_filler.export import to_json, to_markdown
    from note_filler.pipeline import require_traceable_note_product

    src_a = Source(
        id="law:92",
        title="行政程序法第 92 條",
        url=None,
        level="A",
        content="【片段:law:92】行政處分，係指行政機關就公法上具體事件所為之決定。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.1,
    )
    src_b = Source(
        id="law:93",
        title="行政程序法第 93 條",
        url=None,
        level="A",
        content="【片段:law:93】行政機關作成行政處分有裁量權時，得為附款。",
        fetched_date="2026-07-26",
        doc_date=None,
        distance=0.2,
    )
    
    gap_a = Gap("行政處分如何定義？", "missing", "原稿未定義行政處分")
    gap_b = Gap("附款有何限制？", "missing", "原稿未說明附款限制")
    
    claim_a = "行政處分定義參照[^1]。"
    claim_b = "附款限制須兼顧目的[^1]。"
    
    product = assemble_correction(
        _doc(),
        [gap_a, gap_b],
        {gap_a.question: [src_a], gap_b.question: [src_b]},
        {
            gap_a.question: WrittenSupplement(claim_a, [src_a.id]),
            gap_b.question: WrittenSupplement(claim_b, [src_b.id]),
        },
        {
            gap_a.question: cross_validate(gap_a.question, [src_a]),
            gap_b.question: cross_validate(gap_b.question, [src_b]),
        },
    )
    
    # 設定四要素，確保內容互相呼應
    # 論點 A：行政處分定義
    product.segments[1].functional_gap = "原稿未定義行政處分概念，讀者無法理解行政處分的法律定義"
    product.segments[1].user_value = "補齊讀者（行政法初學者）對行政處分定義的理解，避免適用錯誤"
    product.segments[1].related_knowledge = (
        "行政處分定義是適用行政程序法的基礎（支撐決策品質：對應功能缺口「原稿未定義行政處分概念」"
        "提供可追溯依據，降低僅憑印象取捨的風險；"
        "補強使用者理解：幫助讀者正確識別行政處分，避免與其他行政行為混淆）"
    )
    
    # 論點 B：附款限制
    product.segments[2].functional_gap = "原稿未說明附款限制，讀者不知附款不得違背行政處分之目的"
    product.segments[2].user_value = "補齊讀者（行政機關人員）對附款限制的認識，確保行政處分合法性"
    product.segments[2].related_knowledge = (
        "附款限制是保障行政處分合法性的關鍵（支撐決策品質：對應功能缺口「原稿未說明附款限制」"
        "提供可追溯依據，降低僅憑印象取捨的風險；"
        "補強使用者理解：幫助讀者理解附款的目的性限制，避免違法附款）"
    )
    
    require_traceable_note_product(product, source="four-elements-acceptance")
    
    # 落盤成品
    out_dir = Path(tmp_path)
    product_path = out_dir / "note_product.json"
    md_path = out_dir / "note_product.md"
    product_path.write_text(
        json.dumps(to_json(product), ensure_ascii=False, indent=2),
        encoding="utf-8",
        newline="\n",
    )
    md_path.write_text(to_markdown(product), encoding="utf-8", newline="\n")
    report_path = write_binding_report(md_path, product)
    
    # 讀取成品進行驗證
    serialized = json.loads(product_path.read_text(encoding="utf-8"))
    report = parse_binding_report(
        json.loads(report_path.read_text(encoding="utf-8"))
    )
    md = md_path.read_text(encoding="utf-8")
    md_blocks = _parse_md_argument_blocks(md)
    
    assert report["summary"]["all_arguments_ok"] is True
    assert report["summary"]["fail"] == 0
    assert report["argument_count"] == 2
    assert len(md_blocks) == 2
    
    # 逐筆驗證每個論點的四要素
    for arg, block in zip(report["arguments"], md_blocks, strict=True):
        seg = serialized["segments"][arg["segment_index"]]
        assert seg["type"] == "supplement"
        
        argument_id = arg["argument_id"]
        assert argument_id == seg["argument_id"]
        assert argument_id == block["argument_id"]
        
        # 1. 缺口（functional_gap）
        functional_gap = (arg["functional_gap"] or "").strip()
        assert functional_gap, f"{argument_id}：缺口（functional_gap）不可為空"
        assert "原稿未" in functional_gap or "缺失" in functional_gap, (
            f"{argument_id}：缺口應描述原稿缺失"
        )
        assert seg["functional_gap"] == functional_gap
        assert block["functional_gap"] == functional_gap
        
        # 2. 影響對象（從 user_value 中提取）
        user_value = (arg["user_value"] or "").strip()
        assert user_value, f"{argument_id}：使用者價值不可為空"
        # 檢查是否包含受影響的對象（讀者、使用者等）
        has_affected_object = any(
            keyword in user_value 
            for keyword in ["讀者", "使用者", "人員", "當事人", "民眾"]
        )
        assert has_affected_object, (
            f"{argument_id}：使用者價值應描述受影響的對象（讀者、使用者等）"
        )
        assert seg["user_value"] == user_value
        assert block["user_value"] == user_value
        
        # 3. 重要性（從 functional_gap 和 user_value 中提取）
        # 檢查是否包含重要性描述（避免、確保、關鍵、基礎等）
        has_importance = any(
            keyword in functional_gap + user_value
            for keyword in ["避免", "確保", "關鍵", "基礎", "重要", "必要", "核心"]
        )
        assert has_importance, (
            f"{argument_id}：缺口或使用者價值應描述補齊的重要性"
        )
        
        # 4. 關聯知識（related_knowledge）
        related_knowledge = (arg["related_knowledge"] or "").strip()
        assert related_knowledge, f"{argument_id}：關聯知識不可為空"
        assert "支撐決策品質" in related_knowledge, (
            f"{argument_id}：關聯知識必須說明如何支撐決策品質"
        )
        assert "補強使用者理解" in related_knowledge, (
            f"{argument_id}：關聯知識必須說明如何補強使用者理解"
        )
        assert seg["related_knowledge"] == related_knowledge
        assert block["related_knowledge"] == related_knowledge
        
        # 驗證四要素內容互相呼應
        # 提取主題詞
        summary_text = re.sub(r"\[\^\d+\]", "", (arg["argument_text"] or "")).strip()
        
        # 驗證至少在三個要素中有共同的主題詞（更寬鬆的條件）
        all_text = f"{summary_text} {functional_gap} {user_value} {related_knowledge}"
        chinese_tokens = re.findall(r"[\u4e00-\u9fff]{2,}", all_text)
        
        # 找出在三個或更多要素中出現的主題詞
        shared_tokens = []
        for token in chinese_tokens:
            count = sum([
                token in summary_text,
                token in functional_gap,
                token in user_value,
                token in related_knowledge
            ])
            if count >= 3:
                shared_tokens.append(token)
        
        # 如果沒有找到在三個要素中都出現的詞，改為檢查兩兩之間的關聯
        if not shared_tokens:
            # 檢查摘要與缺口之間的共同詞
            summary_gap_shared = [t for t in chinese_tokens if t in summary_text and t in functional_gap]
            # 檢查摘要與使用者價值之間的共同詞
            summary_uv_shared = [t for t in chinese_tokens if t in summary_text and t in user_value]
            # 檢查缺口與關聯知識之間的共同詞
            gap_rk_shared = [t for t in chinese_tokens if t in functional_gap and t in related_knowledge]
            
            # 至少要有兩兩之間的關聯
            assert (summary_gap_shared or summary_uv_shared or gap_rk_shared), (
                f"{argument_id}：四要素應該有兩兩之間的共同主題詞以確保內容互相呼應"
            )
        
        # 驗證關聯知識與缺口、使用者價值的對應關係
        assert functional_gap in related_knowledge or "功能缺口" in related_knowledge, (
            f"{argument_id}：關聯知識應對應功能缺口"
        )
        assert "使用者理解" in related_knowledge or "使用者價值" in related_knowledge, (
            f"{argument_id}：關聯知識應對應使用者價值"
        )
        
        # 驗證來源綁定
        assert arg["binding_ok"] is True
        assert arg["binding_status"] == "pass"
        assert arg["checks"]["has_functional_gap"] is True
        assert arg["checks"]["has_user_value"] is True
        assert arg["checks"]["has_related_knowledge"] is True
        assert arg["checks"]["related_knowledge_consistent"] is True
        assert arg["checks"]["at_least_one_source"] is True
    
    # 驗證 Markdown 摘要可見列包含四要素
    for block in md_blocks:
        visible_summary = block["visible_summary"]
        assert f"argument_id=" in visible_summary
        assert f"functional_gap=" in visible_summary
        assert f"user_value=" in visible_summary
        assert f"related_knowledge=" in visible_summary
        assert "支撐決策品質" in visible_summary
        assert "補強使用者理解" in visible_summary
