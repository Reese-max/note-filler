"""GOAL 0a8e5521d41101e3：恢復「實際筆記」產出 — 非空實際筆記硬性產出閘。

鎖定：
1. 補齊主流程完成後必須有至少一份可追溯、非空 original/supplement 筆記內容
2. 空白或僅稽核摘要（無實質筆記段）必須失敗並帶出原因
3. 不得表面成功（exit/回傳成功）卻無成品

不得弱化既有品質閘：原稿 immutable、pending_evidence、只掛實際引用、法條離線查核。
"""

from __future__ import annotations

import json
import logging
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from note_filler import __main__ as cli
from note_filler.correction import CorrectionDoc, Segment
from note_filler.export import to_markdown
from note_filler.llm import FakeLLM
from note_filler.parse import Document, Paragraph
from note_filler.pipeline import require_non_empty_note_product, run_pipeline
from note_filler.retrieve.models import Source
from tests.test_pipeline import FakeLaw, FakeTwinkle


def _txt(tmp_path: Path, name: str, text: str) -> Path:
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def _docx(tmp_path: Path, name: str, *paragraphs: str) -> Path:
    p = tmp_path / name
    d = DocxDocument()
    for para in paragraphs:
        d.add_paragraph(para)
    d.save(str(p))
    return p


def _src(sid: str = "s1") -> Source:
    return Source(
        id=sid,
        title="來源標題",
        url="https://example.gov.tw/a",
        level="A",
        content="官方結構化記錄全文……",
        fetched_date="2026-07-15",
        doc_date="2026-01-01",
        distance=0.5,
    )


def _full_canned_llm() -> FakeLLM:
    """有原文 + 兩缺口的最小可跑 canned（含補充成品）。"""
    return FakeLLM(
        [
            "admin",
            "正當程序的要件為何?\n聽證程序如何進行?",
            json.dumps(
                [
                    {
                        "question": "正當程序的要件為何?",
                        "status": "missing",
                        "reason": "筆記未展開",
                    },
                    {
                        "question": "聽證程序如何進行?",
                        "status": "missing",
                        "reason": "筆記未提及",
                    },
                ],
                ensure_ascii=False,
            ),
            '{"keyword": "正當程序", "law_name": null}',
            "【待補證】此問題缺乏可用來源,尚待補充。",
            '{"keyword": "聽證", "law_name": null}',
            "聽證程序應保障當事人陳述意見[^1],並依法定程序進行[^2]。",
        ]
    )


def _empty_product_doc(source_path: str = "empty.txt") -> CorrectionDoc:
    """空白成品：無非空 original/supplement（僅空 Document）。"""
    original = Document(source_path=source_path, paragraphs=(), full_text="")
    return CorrectionDoc(original=original, segments=[])


def _whitespace_only_segments_doc() -> CorrectionDoc:
    """僅空白段：等同無實質筆記 / 稽核摘要路徑。"""
    original = Document(
        source_path="ws.txt",
        paragraphs=(),
        full_text="",
    )
    return CorrectionDoc(
        original=original,
        segments=[
            Segment(
                type="original",
                text="   \n\t  ",
                anchor_idx=0,
                sources=[],
                confidence="verified",
                functional_gap="",
                user_value="",
                argument_id="",
            ),
            Segment(
                type="supplement",
                text="",
                anchor_idx=None,
                sources=[],
                confidence="pending_evidence",
                functional_gap="",
                user_value="",
                argument_id="argument:0",
            ),
        ],
    )


# ---- 單元：require_non_empty_note_product ---------------------------------


def test_require_non_empty_rejects_blank_product_with_reason():
    with pytest.raises(RuntimeError, match="未產生非空實際筆記|空白|拒絕視為成功") as ei:
        require_non_empty_note_product(_empty_product_doc("blank-note.txt"))
    msg = str(ei.value)
    assert "空白" in msg or "無任何" in msg


def test_require_non_empty_rejects_whitespace_or_audit_only_with_reason():
    with pytest.raises(RuntimeError, match="未產生非空實際筆記|稽核摘要|拒絕視為成功") as ei:
        require_non_empty_note_product(_whitespace_only_segments_doc())
    assert "稽核摘要" in str(ei.value) or "無實質" in str(ei.value)


def test_require_non_empty_accepts_original_or_supplement():
    original = Document(
        source_path="ok.txt",
        paragraphs=(Paragraph(0, "原文內容"),),
        full_text="原文內容",
    )
    # 僅 original
    require_non_empty_note_product(
        CorrectionDoc(
            original=original,
            segments=[
                Segment(
                    type="original",
                    text="原文內容",
                    anchor_idx=0,
                    sources=[],
                    confidence="verified",
                    functional_gap="",
                    user_value="",
                    argument_id="",
                )
            ],
        )
    )
    # 僅 supplement（無 original 段亦可視為可追溯補充成品）
    require_non_empty_note_product(
        CorrectionDoc(
            original=Document(source_path="s.txt", paragraphs=(), full_text=""),
            segments=[
                Segment(
                    type="supplement",
                    text="可追溯補充內容",
                    anchor_idx=None,
                    sources=[],
                    confidence="pending_evidence",
                    functional_gap="",
                    user_value="",
                    argument_id="argument:0",
                )
            ],
        )
    )


def test_require_non_empty_emits_audit_event(caplog):
    with caplog.at_level(logging.WARNING, logger="note_filler.pipeline"):
        with pytest.raises(RuntimeError, match="未產生非空實際筆記"):
            require_non_empty_note_product(_empty_product_doc("audit-id.txt"))
    assert "note_product_empty" in caplog.text
    assert "audit-id.txt" in caplog.text


# ---- 主流程 run_pipeline -------------------------------------------------


def test_run_pipeline_produces_non_empty_traceable_notes(tmp_path):
    note = _docx(
        tmp_path,
        "note.docx",
        "行政程序法要求行政行為應遵守正當程序。",
        "本筆記僅記錄部分重點,尚未展開。",
    )
    llm = _full_canned_llm()
    twinkle = FakeTwinkle(
        [
            [],
            [_src("s1"), _src("s2")],
        ]
    )
    doc = run_pipeline(str(note), llm, twinkle, FakeLaw())
    texts = [s.text.strip() for s in doc.segments if (s.text or "").strip()]
    assert texts, "主流程必須產出至少一份非空可追溯筆記內容"
    originals = [s for s in doc.segments if s.type == "original" and s.text.strip()]
    supplements = [s for s in doc.segments if s.type == "supplement" and s.text.strip()]
    assert originals, "原文段應保留"
    assert supplements, "應有補充成品"


def test_main_flow_actual_note_has_traceable_source_and_extension(tmp_path):
    """最小回歸：直接跑主流程，產出非空白實際筆記，且同時含可追溯來源與延伸論點。

    鎖定失敗語義（精準失敗以鎖定缺口）：
    - 空輸出 / 僅空白段
    - 僅測試稽核或【待補證】占位、無實質延伸論點
    - 有文字但無可追溯來源（無 sources、匯出無 [^n]/參考列）

    不改弱既有品質閘：原稿 immutable、pending_evidence、只掛實際引用、法條離線查核。
    離線 FakeLLM/FakeTwinkle；非 integration。
    """
    note = _docx(
        tmp_path,
        "note.docx",
        "行政程序法要求行政行為應遵守正當程序。",
        "本筆記僅記錄部分重點,尚未展開。",
    )
    llm = _full_canned_llm()
    twinkle = FakeTwinkle(
        [
            [],
            [_src("s1"), _src("s2")],
        ]
    )
    doc = run_pipeline(str(note), llm, twinkle, FakeLaw())
    md = to_markdown(doc)

    assert isinstance(md, str) and md.strip(), (
        "主流程回空輸出：未產出非空白實際筆記；拒絕視為成功"
    )
    assert "note_product_empty" not in md, (
        "輸出疑似測試稽核事件字串而非實際筆記"
    )

    note_segments = [
        s
        for s in doc.segments
        if s.type in ("original", "supplement") and (s.text or "").strip()
    ]
    assert note_segments, (
        "主流程僅有空段/稽核狀態、無非空白實際筆記內容"
    )

    # 可追溯來源：至少一補充段掛實際 sources，且匯出可對到 [^n] 或參考列
    traced = [
        s
        for s in doc.segments
        if s.type == "supplement" and s.sources and (s.text or "").strip()
    ]
    assert traced, (
        "實際筆記缺少可追溯來源：無任何掛 sources 的補充段"
        "（若流程只回待補證/稽核則在此精準失敗）"
    )
    for seg in traced:
        for src in seg.sources:
            assert src.id, f"來源缺 id，無法追溯: {src!r}"
            assert src.title or src.url, f"來源缺 title/url: {src!r}"
    assert "[^" in md, (
        "匯出 markdown 缺少 [^n] 註腳/參考列，來源不可從成品追溯"
    )

    # 至少一個延伸論點：實質補充文字，非僅【待補證】占位
    extensions = [
        s
        for s in doc.segments
        if s.type == "supplement"
        and (s.text or "").strip()
        and not (s.text or "").strip().startswith("【待補證】")
    ]
    assert extensions, (
        "實際筆記缺少延伸論點：補充段全為【待補證】/稽核占位、無實質延伸"
    )
    assert any(len(s.text.strip()) >= 10 for s in extensions), (
        "延伸論點內容過短，疑似非實質補充"
    )
    # 延伸論點應出現在匯出正文（非僅內部狀態）
    assert any(ext.text.strip()[:20] in md for ext in extensions), (
        "延伸論點未進入匯出成品（僅內部狀態/稽核）"
    )


def test_run_pipeline_empty_input_and_no_gaps_fails_product_gate(tmp_path, caplog):
    """空輸入 + 問題生成空清單 → 無段 → 不得表面成功。"""
    note = _txt(tmp_path, "empty.txt", "   \n\n  ")
    # domain / questions(JSON wrapper → []) → 無 gap → 空 segments
    llm = FakeLLM(["other", '{"questions": []}'])
    with caplog.at_level(logging.WARNING, logger="note_filler.pipeline"):
        with pytest.raises(RuntimeError, match="未產生非空實際筆記|拒絕視為成功"):
            run_pipeline(str(note), llm, FakeTwinkle([]), FakeLaw())
    assert "note_product_empty" in caplog.text


# ---- CLI process_file 交付層 ---------------------------------------------


def test_process_file_rejects_empty_product_stub_with_reason(tmp_path, monkeypatch):
    note = _txt(tmp_path, "note.txt", "有輸入但 pipeline 回空成品")
    monkeypatch.setattr(cli, "run_pipeline", lambda *a, **k: _empty_product_doc(str(note)))
    with pytest.raises(RuntimeError, match="未產生非空實際筆記|空白|拒絕視為成功"):
        cli.process_file(note, None, None, None, out_dir=None, fmt="md")
    assert not (tmp_path / "note.訂正稿.md").exists()


def test_process_file_success_requires_non_empty_body(tmp_path, monkeypatch):
    note = _txt(tmp_path, "note.txt", "一、標題\n內容")

    class _Ok:
        segments = [
            Segment(
                type="original",
                text="一、標題\n內容",
                anchor_idx=0,
                sources=[],
                confidence="verified",
                functional_gap="",
                user_value="",
                argument_id="",
            ),
            Segment(
                type="supplement",
                text="補充說明",
                anchor_idx=0,
                sources=[],
                confidence="pending_evidence",
                functional_gap="",
                user_value="",
                argument_id="argument:0",
            ),
        ]

    monkeypatch.setattr(cli, "run_pipeline", lambda *a, **k: _Ok())
    monkeypatch.setattr(cli, "to_markdown", lambda doc: "# 訂正稿\n一、標題\n內容\n\n> 【補充】補充說明")

    r = cli.process_file(note, None, None, None, out_dir=None, fmt="md")
    dest = Path(r["output"])
    assert dest.exists()
    body = dest.read_text(encoding="utf-8")
    assert body.strip()
    assert "一、標題" in body or "補充" in body
