"""PR #14: review decisions must stay evidence-bound and durable."""

from dataclasses import replace
import json
from pathlib import Path

import pytest
import httpx
from docx import Document as DocxDocument

import app.server as server
from note_filler.export import to_docx, to_json, to_markdown
from note_filler import __main__ as cli
from note_filler.parse import Paragraph
from tests.test_review_queue import _doc, _ledger
from tests.review_forms import review_form, seed_result


@pytest.fixture(autouse=True)
def require_review_api():
    # Keep a missing feature red at test setup, rather than abort collection.
    global ReviewLedger, ReviewState, source_stances
    from note_filler.review import ReviewLedger, ReviewState, source_stances


@pytest.mark.parametrize("clear_ids", [False, True])
def test_reaccepting_missing_evidence_stays_stale(clear_ids):
    doc = _doc()
    seg = doc.segments[1]
    ledger = _ledger(doc)
    ledger.record(seg, "accepted")
    seg.sources.clear()
    if clear_ids:
        seg.source_ids.clear()
    ledger.record(seg, "accepted")

    detail = ledger.state_detail(seg)
    assert detail["state"] == ReviewState.STALE_REVIEW
    assert detail["stale_reason"] == "evidence_unavailable"
    assert seg.text not in to_markdown(doc, export_mode="accepted-only", ledger=ledger)


@pytest.mark.parametrize("unavailable", ["citation_only", "empty_fragment"])
def test_reaccepting_unavailable_source_fragments_stays_stale(unavailable):
    doc = _doc()
    seg = doc.segments[1]
    ledger = _ledger(doc)
    ledger.record(seg, "accepted")
    if unavailable == "citation_only":
        seg.citation_spans[0]["source_id"] = "missing"
    else:
        seg.sources[0].content = " \n\t"
    ledger.record(seg, "accepted")
    assert ledger.state_detail(seg)["stale_reason"] == "evidence_unavailable"
    assert seg.text not in to_markdown(doc, export_mode="accepted-only", ledger=ledger)
    missing = [s for s in source_stances(seg) if s["missing"]]
    assert missing and all(s["stance"] == "unresolved" for s in missing)


@pytest.mark.parametrize("conflict_note", ["來源表述不一致", None])
def test_source_stances_are_relative_to_negative_claim(conflict_note):
    seg = _doc().segments[3]
    seg.text = "沒收處分不應以判決為準。"
    seg.conflict_note = conflict_note
    assert {s["source_id"]: s["stance"] for s in source_stances(seg)} == {
        "s3": "conflicts",
        "s4": "supports",
    }


@pytest.mark.parametrize("export_mode", ["review-draft", "accepted-only"])
@pytest.mark.parametrize("fmt", ["md", "json", "docx"])
def test_export_does_not_replay_a_different_documents_ledger(tmp_path, fmt, export_mode):
    doc = _doc()
    ledger = _ledger(doc)
    ledger.record(doc.segments[1], "accepted")
    other = replace(doc, original=replace(
        doc.original, full_text="另一份原稿。",
        paragraphs=(Paragraph(idx=0, text="另一份原稿。"),),
    ))
    kwargs = {"export_mode": export_mode, "ledger": ledger}
    if fmt == "json":
        payload = to_json(other, **kwargs)
        if export_mode == "review-draft":
            assert payload["segments"][1]["review_state"] == "unreviewed"
        else:
            assert all(s["type"] == "original" for s in payload["segments"])
    elif fmt == "md":
        text = to_markdown(other, **kwargs)
        if export_mode == "review-draft":
            assert "review_state=accepted" not in text
        else:
            assert doc.segments[1].text not in text
    else:
        path = tmp_path / "export.docx"
        to_docx(other, str(path), **kwargs)
        text = "\n".join(p.text for p in DocxDocument(path).paragraphs)
        if export_mode == "accepted-only":
            assert doc.segments[1].text not in text
        else:
            assert "review_state=accepted" not in text


def test_docx_review_draft_marks_rejected_claim_and_mode(tmp_path):
    doc = _doc()
    ledger = _ledger(doc)
    ledger.record(doc.segments[1], "rejected")
    path = tmp_path / "draft.docx"
    to_docx(doc, str(path), export_mode="review-draft", ledger=ledger)
    exported = DocxDocument(path)
    text = "\n".join(p.text for p in exported.paragraphs)
    header = "\n".join(p.text for p in exported.sections[0].header.paragraphs)
    assert doc.segments[1].text in text
    assert "review_state=rejected" in text
    assert "export_mode=review-draft" in text + header


@pytest.mark.anyio
async def test_edited_accepted_has_a_queue_filter(async_client, monkeypatch, tmp_path):
    doc = _doc()
    ledger = _ledger(doc)
    ledger.record(doc.segments[1], "edited_accepted")
    result_id = seed_result(doc, ledger)
    monkeypatch.setattr(server, "LEDGER_PATH", tmp_path / "ledger.json")
    response = await async_client.get(f"/result/{result_id}")
    assert f'href="/result/{result_id}?filter=edited_accepted"' in response.text
    response = await async_client.get(f"/result/{result_id}?filter=edited_accepted")
    assert 'data-review-state="edited_accepted"' in response.text
    assert 'data-review-state="unreviewed"' not in response.text


@pytest.mark.anyio
async def test_next_pending_leaves_a_filter_that_hides_the_claim(async_client, monkeypatch, tmp_path):
    doc = _doc()
    ledger = _ledger(doc)
    ledger.record(doc.segments[1], "accepted")
    result_id = seed_result(doc, ledger)
    monkeypatch.setattr(server, "LEDGER_PATH", tmp_path / "ledger.json")
    response = await async_client.get(f"/result/{result_id}?filter=accepted")
    assert 'id="claim-argument:1"' not in response.text
    assert f'href="/result/{result_id}#claim-argument:1"' in response.text


@pytest.mark.parametrize("edited_text", ["", "尚未存檔的人工修訂。"])
@pytest.mark.anyio
async def test_failed_save_does_not_publish_review_or_edit(monkeypatch, tmp_path, edited_text):
    doc = _doc()
    seg = doc.segments[1]
    before_text = seg.text
    ledger = _ledger(doc)
    ledger.record(seg, "rejected")
    result_id = seed_result(doc, ledger)
    monkeypatch.setattr(server, "LEDGER_PATH", tmp_path / "ledger.json")
    ledger.save(server._ledger_path_for(doc, result_id))
    before_ledger = ledger.to_dict()

    def fail_save(self, path):
        raise OSError("disk full")

    monkeypatch.setattr(ReviewLedger, "save", fail_save)
    transport = httpx.ASGITransport(app=server.app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        response = await client.post(f"/review/{result_id}", data={
            **await review_form(client, seg.argument_id, result_id=result_id),
            "argument_id": seg.argument_id,
            "decision": "accepted",
            "edited_text": edited_text,
        })
        assert response.status_code == 500
        assert server.app.state.results[result_id]["doc"].segments[1].text == before_text
        assert server.app.state.results[result_id]["review_ledger"].to_dict() == before_ledger
        assert ReviewLedger.load(server._ledger_path_for(doc, result_id)).to_dict() == before_ledger
        assert "行政處分係指" not in (await client.get(f"/export/{result_id}?mode=accepted-only")).text

    assert doc.original.full_text == doc.segments[0].text


@pytest.mark.parametrize("new_upload", [False, True])
@pytest.mark.anyio
async def test_review_does_not_overwrite_a_document_changed_while_reading_form(
    async_client, monkeypatch, tmp_path, new_upload,
):
    doc = _doc()
    ledger = _ledger(doc)
    concurrent = replace(doc, segments=[
        replace(s, text="較新修訂。") if s.type == "supplement" else s
        for s in doc.segments
    ])
    if new_upload:
        concurrent.original = replace(doc.original, full_text="新上傳。")
    concurrent_ledger = _ledger(concurrent)
    concurrent_ledger.record(concurrent.segments[1], "edited_accepted")
    result_id = seed_result(doc, ledger)
    monkeypatch.setattr(server, "LEDGER_PATH", tmp_path / "ledger.json")
    form = await review_form(async_client, result_id=result_id)
    parse_form = server.Request.form

    async def interleaved_form(request):
        form = await parse_form(request)
        concurrent_ledger.save(server._ledger_path_for(concurrent, result_id))
        server.app.state.results[result_id]["doc"] = concurrent
        server.app.state.results[result_id]["generated_doc"] = concurrent
        server.app.state.results[result_id]["review_ledger"] = concurrent_ledger
        return form

    monkeypatch.setattr(server.Request, "form", interleaved_form)
    response = await async_client.post(f"/review/{result_id}", data={
        **form,
        "argument_id": "argument:0", "decision": "rejected",
    })
    assert response.status_code == 409
    assert server.app.state.results[result_id]["doc"] is concurrent
    assert server.app.state.results[result_id]["review_ledger"] is concurrent_ledger
    assert concurrent_ledger.state_of(concurrent.segments[1]) == ReviewState.EDITED_ACCEPTED


@pytest.mark.parametrize("fmt", ["json", "md", "docx"])
def test_accepted_only_metrics_describe_exported_claims(tmp_path, fmt):
    doc = _doc()
    ledger = _ledger(doc)
    ledger.record(doc.segments[1], "accepted")
    kwargs = {"export_mode": "accepted-only", "ledger": ledger}
    if fmt == "json":
        exported = to_json(doc, **kwargs)
        assert exported["binding_summary"]["argument_count"] == 1
        assert set(exported["polaris_metrics"]["claim_source_map"]) == {"argument:0"}
        text = json.dumps(exported)
    elif fmt == "md":
        text = to_markdown(doc, **kwargs)
    else:
        path = tmp_path / "accepted.docx"
        to_docx(doc, str(path), **kwargs)
        text = "\n".join(p.text for p in DocxDocument(path).paragraphs)
    assert "argument:1" not in text and "argument:2" not in text
    assert len(doc.segments) == 4


def test_cli_receipt_counts_only_delivered_claims(tmp_path, monkeypatch):
    from note_filler.pipeline import run_pipeline
    from tests.test_output_consistency import _note_fixture, _pipeline_canned

    note = Path(_note_fixture(tmp_path, "行政程序法要求行政行為應遵守正當程序。"))
    doc = run_pipeline(str(note), *_pipeline_canned())
    accepted = next(s for s in doc.segments if s.type == "supplement" and s.confidence == "verified")
    ledger = _ledger(doc)
    ledger.record(accepted, "accepted")
    path = ledger.save(tmp_path / "ledger.json")
    monkeypatch.setattr(cli, "run_pipeline", lambda *args: doc)
    result = cli.process_file(
        note, None, None, None, tmp_path / "out", "json",
        state_dir=tmp_path / "state", export_mode="accepted-only", ledger_path=path,
    )
    assert result["supplements"] == result["verified"] == 1
    manifest = json.loads((tmp_path / "out" / "delivery_manifest.json").read_text(encoding="utf-8"))
    report = json.loads((tmp_path / "out" / "binding_report.json").read_text(encoding="utf-8"))
    assert manifest["supplements"] == report["argument_count"] == 1
    assert set(manifest["polaris_metrics"]["claim_source_map"]) == {accepted.argument_id}
