import json
import socket
import pytest
from docx import Document as DocxDocument

from note_filler.llm import FakeLLM
from note_filler.retrieve.models import Source          # Source 定義處(T6/型別鎖定)
from note_filler.pipeline import run_pipeline


def _grok_reachable(host: str = "127.0.0.1", port: int = 8318) -> bool:
    try:
        with socket.create_connection((host, port), timeout=1.0):
            return True
    except OSError:
        return False


@pytest.fixture
def note_path(tmp_path):
    """用 python-docx 現造一份真 .docx,避免依賴外部檔。"""
    p = tmp_path / "note.docx"
    d = DocxDocument()
    d.add_paragraph("行政程序法要求行政行為應遵守正當程序。")
    d.add_paragraph("本筆記僅記錄部分重點,尚未展開。")
    d.save(str(p))
    return str(p)


class FakeTwinkle:
    """依序回傳每個 gap 的來源批次;不打真網路。"""
    def __init__(self, batches):
        self.batches = list(batches)
        self.queries = []

    def search(self, query, n=3):
        self.queries.append(query)
        return self.batches.pop(0) if self.batches else []


class FakeLaw:
    def lookup_article(self, law_name, article_no): return None
    def law_exists(self, name): return True
    def fuzzy_find_law(self, name): return None
    def search_articles(self, keyword, limit=5, law_name=None): return []


def _src(sid, title, url, level):
    return Source(
        id=sid, title=title, url=url, level=level,
        content=f"{title} 官方結構化記錄全文……",
        fetched_date="2026-07-15", doc_date="2026-01-01", distance=0.6,
    )


def test_run_pipeline_invariant(note_path):
    # admin 為法條領域:每 gap 呼叫序=(keyword 抽取 JSON + writer)。canned 依序:
    # domain / questions / gaps / gap1(keyword,writer) / gap2(keyword,writer) 共 7 個。
    llm = FakeLLM([
        "admin",                                              # detect_domain 取單一標籤
        "正當程序的要件為何?\n聽證程序如何進行?",              # generate_questions 逐行解析
        json.dumps([                                          # detect_gaps 一次回陣列
            {"question": "正當程序的要件為何?", "status": "missing", "reason": "筆記未展開"},
            {"question": "聽證程序如何進行?", "status": "missing", "reason": "筆記未提及"},
        ], ensure_ascii=False),
        '{"keyword": "正當程序", "law_name": null}',           # gap1 法條關鍵詞抽取(FakeLaw 回 [])
        "【待補證】此問題缺乏可用來源,尚待補充。",              # gap1 無源 → writer 待補證
        '{"keyword": "聽證", "law_name": null}',               # gap2 法條關鍵詞抽取(FakeLaw 回 [])
        "聽證程序應保障當事人陳述意見[^1],並依法定程序進行[^2]。",  # gap2 引用兩獨立源
    ])
    twinkle = FakeTwinkle([
        [],                                                   # gap1 無來源 → pending_evidence
        [_src("s1", "行政院公報", "https://a", "A"),
         _src("s2", "立法院議案", "https://b", "B")],          # gap2 兩獨立 A/B → verified
    ])

    doc = run_pipeline(note_path, llm, twinkle, FakeLaw())

    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應至少有一個補充段"

    # C6 不變式:斷言「每個無源 supplement 的 confidence==pending_evidence」,
    # 不斷言「pending 一定存在」;有 >=2 獨立 A/B 源則為 verified。
    for seg in supplements:
        if not seg.sources:
            assert seg.confidence == "pending_evidence"
        elif len(seg.sources) >= 2:
            assert seg.confidence == "verified"

    # supplement text 是寫出來的句子(非空);gap2 引用兩獨立 A/B 源 → verified 且掛 2 來源
    gap2 = [s for s in supplements if s.sources]
    assert gap2 and gap2[0].confidence == "verified"
    assert len(gap2[0].sources) == 2 and "聽證程序" in gap2[0].text
    # 每個 gap 恰觸發一次 twinkle retrieve
    assert len(twinkle.queries) == 2


def test_run_pipeline_law_domain_runs_citation_check(note_path, monkeypatch):
    import note_filler.pipeline as pl

    calls = []
    def fake_check(text, lookup):        # 對齊 C2:關鍵字 text=... 會綁到此參數
        calls.append(text)
        return []                        # 無問題引用 → 不降級
    monkeypatch.setattr(pl, "check_law_citations", fake_check)

    llm = FakeLLM([
        "law",                                                   # domain 標籤
        "民法第184條的構成要件為何?",                             # 換行問題(單行)
        json.dumps([{"question": "民法第184條的構成要件為何?",
                     "status": "missing", "reason": "缺"}], ensure_ascii=False),
        '{"keyword": "侵權行為", "law_name": null}',              # gap 法條關鍵詞抽取(FakeLaw 回 [])
        "侵權行為以故意或過失不法侵害他人權利為要件[^1][^2]。",    # writer 撰寫補充
    ])
    twinkle = FakeTwinkle([[_src("s1", "法規原文A", "https://a", "A"),
                           _src("s2", "立法院議案B", "https://b", "B")]])

    doc = run_pipeline(note_path, llm, twinkle, FakeLaw())

    supplements = [s for s in doc.segments if s.type == "supplement"]
    # law 領域:每個補充段都經過 check_law_citations(text=...)
    assert len(calls) == len(supplements)
    assert calls, "law 領域至少應跑一次法規引用檢查"


def test_run_pipeline_malformed_gap_output_falls_back_to_pending(note_path):
    llm = FakeLLM([
        "law",
        "正當程序的要件為何?",
        "這不是 JSON",
        '{"keyword": "不存在", "law_name": null}',
        "【待補證】模型回應無法解析,尚待補充。",
    ])
    twinkle = FakeTwinkle([[]])

    doc = run_pipeline(note_path, llm, twinkle, FakeLaw())

    originals = [s.text for s in doc.segments if s.type == "original"]
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert originals == [
        "行政程序法要求行政行為應遵守正當程序。",
        "本筆記僅記錄部分重點,尚未展開。",
    ]
    assert len(supplements) == 1
    assert supplements[0].text.startswith("【待補證】")
    assert supplements[0].confidence == "pending_evidence"
    assert supplements[0].sources == []
    assert twinkle.queries == ["正當程序的要件為何?"]


@pytest.mark.integration
@pytest.mark.skipif(not _grok_reachable(), reason="grok proxy(127.0.0.1:8318)未上線,條件式略過")
def test_run_pipeline_real_grok(note_path):
    """打真 grok(http://127.0.0.1:8318/v1, grok-4.3);twinkle/law 用 fake 隔離,
    驗 parse→domain→questions→gaps→assemble 整條在真模型輸出下不炸。
    """
    from note_filler.llm import GrokClient

    llm = GrokClient()                        # base_url/model 依鎖定預設
    twinkle = FakeTwinkle([[], [], [], []])   # 每個 gap 都回空,補充段一律 pending_evidence
    doc = run_pipeline(note_path, llm, twinkle, FakeLaw())

    assert doc.original is not None
    assert isinstance(doc.segments, list)
    # 真模型下無源補充仍須守 C6 不變式
    for seg in doc.segments:
        if seg.type == "supplement" and not seg.sources:
            assert seg.confidence == "pending_evidence"
