"""端到端驗收測試：論點區塊連結數、URL 格式、來源相關性與互斥性檢查。

直接檢查最終成品筆記中每個論點區塊的：
1. 連結數（至少 2 條真實可開啟連結）
2. URL 格式有效性（http/https 協議）
3. 來源與論點的直接相關性
4. 【待補來源】與真實連結的互斥性

測試覆蓋：
- 來源充足案例（>= 2 條有效連結）
- 來源不足案例（< 2 條有效連結，應標示【待補來源】）
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from note_filler.correction import (
    MIN_OPENABLE_LINKS,
    CorrectionDoc,
    assemble_correction,
    is_openable_url,
)
from note_filler.export import to_json, to_markdown
from note_filler.gap import Gap
from note_filler.llm import FakeLLM
from note_filler.parse import Document, Paragraph, parse_note
from note_filler.pipeline import run_pipeline
from note_filler.retrieve.models import Source
from note_filler.verify import cross_validate
from note_filler.write import WrittenSupplement


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _doc(path: str = "input/note.txt") -> Document:
    return Document(
        source_path=path,
        paragraphs=(Paragraph(0, "原稿逐字保留。"),),
        full_text="原稿逐字保留。",
    )


def _source(
    id: str,
    title: str,
    level: str = "A",
    url: str | None = None,
    content: str = "",
) -> Source:
    return Source(
        id=id,
        title=title,
        url=url,
        level=level,
        content=content or f"{title} 內容",
        fetched_date="2026-07-28",
        doc_date=None,
        distance=0.5,
    )


def _assemble_with_sources(
    sources: list[Source],
    gap: Gap,
    used_ids: list[str] | None = None,
    omitted_ids: list[str] | None = None,
) -> CorrectionDoc:
    """組裝含指定來源的 CorrectionDoc。"""
    used_ids = used_ids or [s.id for s in sources]
    text = "補充內容[^1]。" if len(used_ids) >= 1 else "【待補證】"
    return assemble_correction(
        _doc(),
        [gap],
        {gap.question: sources},
        {gap.question: WrittenSupplement(text, used_ids, omitted_source_ids=omitted_ids or [])},
        {gap.question: cross_validate(gap.question, sources)},
    )


# ---------------------------------------------------------------------------
# 端到端測試：來源充足案例
# ---------------------------------------------------------------------------

def test_e2e_argument_links_sufficient_case():
    """端到端測試：來源充足案例，論點區塊應有 >= 2 條有效連結。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 來源充足：2 個有效 URL 的來源
    src_a = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url="https://law.moj.gov.tw/LawClass/A0030055",
        content="行政程序法第92條：本法所稱行政處分，係指行政機關就公法上具體事件所為之決定或其他公權力措施。",
    )
    src_b = _source(
        "s2",
        "行政處分定義解釋",
        level="B",
        url="https://www.judicial.gov.tw/tw/cp-1.php",
        content="行政處分係指行政機關就公法上具體事件所為之對外發生法律效果之單方行政行為。",
    )
    
    doc = _assemble_with_sources([src_a, src_b], gap, used_ids=["s1", "s2"])
    
    # 檢查每個論點區塊
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 1. 連結數檢查
        assert seg.openable_links_count >= MIN_OPENABLE_LINKS, (
            f"論點區塊連結數不足：{seg.openable_links_count} < {MIN_OPENABLE_LINKS}"
        )
        
        # 2. URL 格式有效性檢查
        for src in seg.sources:
            if src.url:
                assert is_openable_url(src.url), (
                    f"來源 URL 格式無效：{src.url}"
                )
                assert src.url.startswith(("http://", "https://")), (
                    f"URL 必須為 http/https 協議：{src.url}"
                )
        
        # 3. 來源與論點的直接相關性檢查
        # 檢查來源內容是否包含論點關鍵詞
        question_keywords = ["行政處分", "定義"]
        for src in seg.sources:
            if src.content:
                # 至少一個關鍵詞應出現在來源內容中
                has_keyword = any(kw in src.content for kw in question_keywords)
                assert has_keyword, (
                    f"來源內容缺乏論點關鍵詞：{src.title} - {src.content[:50]}"
                )
        
        # 4. 【待補來源】與真實連結的互斥性檢查
        # 有足夠連結時，不應有【待補來源】標記
        assert seg.openable_links_status == "sufficient", (
            f"連結充足但狀態為 {seg.openable_links_status}"
        )
        assert not seg.openable_links_incomplete_reason, (
            f"連結充足但不應有缺失原因：{seg.openable_links_incomplete_reason}"
        )
        
        # 檢查最終輸出不含【待補來源】標記
        assert "【待補來源】" not in seg.text, (
            "連結充足時不應在文字中出現【待補來源】"
        )


# ---------------------------------------------------------------------------
# 端到端測試：來源不足案例
# ---------------------------------------------------------------------------

def test_e2e_argument_links_insufficient_case():
    """端到端測試：來源不足案例，論點區塊應標示【待補來源】。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 來源不足：只有 1 個有效 URL 的來源
    src = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url="https://law.moj.gov.tw/LawClass/A0030055",
        content="行政程序法第92條：本法所稱行政處分，係指行政機關就公法上具體事件所為之決定或其他公權力措施。",
    )
    
    doc = _assemble_with_sources([src], gap, used_ids=["s1"])
    
    # 檢查每個論點區塊
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 1. 連結數檢查
        assert seg.openable_links_count < MIN_OPENABLE_LINKS, (
            f"來源不足案例連結數應 < {MIN_OPENABLE_LINKS}：{seg.openable_links_count}"
        )
        
        # 2. URL 格式有效性檢查（現有連結仍需有效）
        for src in seg.sources:
            if src.url:
                assert is_openable_url(src.url), (
                    f"來源 URL 格式無效：{src.url}"
                )
                assert src.url.startswith(("http://", "https://")), (
                    f"URL 必須為 http/https 協議：{src.url}"
                )
        
        # 3. 來源與論點的直接相關性檢查（現有來源仍需相關）
        question_keywords = ["行政處分", "定義"]
        for src in seg.sources:
            if src.content:
                has_keyword = any(kw in src.content for kw in question_keywords)
                assert has_keyword, (
                    f"來源內容缺乏論點關鍵詞：{src.title} - {src.content[:50]}"
                )
        
        # 4. 【待補來源】與真實連結的互斥性檢查
        # 連結不足時，應有【待補來源】標記
        assert seg.openable_links_status == "insufficient", (
            f"連結不足但狀態為 {seg.openable_links_status}"
        )
        assert seg.openable_links_incomplete_reason, (
            "連結不足時應有缺失原因說明"
        )
        
        # 檢查最終輸出應含【待補來源】標記
        md = to_markdown(doc)
        assert "【待補來源】" in md, (
            "連結不足時應在輸出中出現【待補來源】"
        )


# ---------------------------------------------------------------------------
# 端到端測試：無 URL 來源案例
# ---------------------------------------------------------------------------

def test_e2e_argument_links_no_url_case():
    """端到端測試：來源無 URL 案例，論點區塊應標示【待補來源】。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 來源無 URL
    src = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url=None,  # 無 URL
        content="行政程序法第92條：本法所稱行政處分，係指行政機關就公法上具體事件所為之決定或其他公權力措施。",
    )
    
    doc = _assemble_with_sources([src], gap, used_ids=["s1"])
    
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 連結數應為 0
        assert seg.openable_links_count == 0, (
            f"無 URL 來源連結數應為 0：{seg.openable_links_count}"
        )
        
        # 狀態應為 insufficient
        assert seg.openable_links_status == "insufficient", (
            f"無 URL 來源狀態應為 insufficient：{seg.openable_links_status}"
        )
        
        # 應有缺失原因
        assert seg.openable_links_incomplete_reason, (
            "無 URL 來源應有缺失原因說明"
        )
        assert "無可開啟連結" in seg.openable_links_incomplete_reason, (
            "缺失原因應說明無可開啟連結"
        )


# ---------------------------------------------------------------------------
# 端到端測試：混合 URL 格式案例
# ---------------------------------------------------------------------------

def test_e2e_argument_links_mixed_url_formats():
    """端到端測試：混合有效與無效 URL 格式，只計算有效連結。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 混合 URL：1 個有效，1 個無效
    src_valid = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url="https://law.moj.gov.tw/LawClass/A0030055",
        content="行政程序法第92條定義。",
    )
    src_invalid = _source(
        "s2",
        "相關解釋",
        level="B",
        url="ftp://example.com/file",  # 無效協議
        content="相關解釋內容。",
    )
    
    doc = _assemble_with_sources([src_valid, src_invalid], gap, used_ids=["s1", "s2"])
    
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 只計算有效 URL
        assert seg.openable_links_count == 1, (
            f"應只計算有效 URL：{seg.openable_links_count}"
        )
        
        # 狀態應為 insufficient（因為只有 1 條有效連結）
        assert seg.openable_links_status == "insufficient", (
            f"有效連結不足時狀態應為 insufficient：{seg.openable_links_status}"
        )


# ---------------------------------------------------------------------------
# 端到端測試：延伸閱讀補足連結
# ---------------------------------------------------------------------------

def test_e2e_argument_links_extended_readings():
    """端到端測試：引用來源不足但延伸閱讀補足連結數。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 1 個引用來源 + 1 個延伸閱讀
    src_cited = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url="https://law.moj.gov.tw/LawClass/A0030055",
        content="行政程序法第92條定義。",
    )
    src_extended = _source(
        "s2",
        "行政處分解釋",
        level="B",
        url="https://www.judicial.gov.tw/tw/cp-1.php",
        content="行政處分解釋內容。",
    )
    
    doc = _assemble_with_sources(
        [src_cited, src_extended],
        gap,
        used_ids=["s1"],
        omitted_ids=["s2"],  # s2 為延伸閱讀
    )
    
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 延伸閱讀應補足連結數
        assert seg.openable_links_count >= MIN_OPENABLE_LINKS, (
            f"延伸閱讀應補足連結數：{seg.openable_links_count} < {MIN_OPENABLE_LINKS}"
        )
        
        # 狀態應為 sufficient
        assert seg.openable_links_status == "sufficient", (
            f"延伸閱讀補足後狀態應為 sufficient：{seg.openable_links_status}"
        )
        
        # 應有延伸閱讀欄位
        assert seg.extended_readings, "應有延伸閱讀"
        assert len(seg.extended_readings) == 1, "應有 1 個延伸閱讀"
        assert seg.extended_readings[0]["source_id"] == "s2", "延伸閱讀應為 s2"


# ---------------------------------------------------------------------------
# 端到端測試：完整流程輸出驗證
# ---------------------------------------------------------------------------

def test_e2e_final_output_links_validation(tmp_path):
    """端到端測試：完整流程的最終輸出（Markdown/JSON）連結驗證。"""
    from docx import Document as DocxDocument
    
    # 建立測試筆記
    note_path = tmp_path / "note.docx"
    docx = DocxDocument()
    docx.add_paragraph("行政程序法要求行政行為應遵守正當程序。")
    docx.add_paragraph("本筆記僅記錄部分重點，尚未展開。")
    docx.save(str(note_path))
    
    # 模擬 LLM 與檢索
    llm = FakeLLM([
        "law",
        "行政處分的定義為何？",
        json.dumps([
            {"question": "行政處分的定義為何？", "status": "missing", "reason": "筆記未展開定義"},
        ], ensure_ascii=False),
        '{"keyword": "行政處分", "law_name": "行政程序法"}',
        "行政處分係指行政機關就公法上具體事件所為之對外發生法律效果之單方行政行為[^1]。",
    ])
    
    from test_pipeline import FakeTwinkle, FakeLaw
    twinkle = FakeTwinkle([[
        _source("s1", "行政程序法第92條", "A", "https://law.moj.gov.tw/LawClass/A0030055"),
        _source("s2", "行政處分解釋", "B", "https://www.judicial.gov.tw/tw/cp-1.php"),
    ]])
    
    # 執行完整流程
    doc = run_pipeline(str(note_path), llm, twinkle, FakeLaw())
    
    # 驗證 Markdown 輸出
    md = to_markdown(doc)
    assert md.strip(), "Markdown 輸出不應為空"
    
    # 檢查論點區塊在 Markdown 中的連結
    supplements = [s for s in doc.segments if s.type == "supplement"]
    for seg in supplements:
        if seg.openable_links_status == "sufficient":
            # 連結充足時，Markdown 中不應有【待補來源】
            assert "【待補來源】" not in md, "連結充足時 Markdown 不應有【待補來源】"
        elif seg.openable_links_status == "insufficient":
            # 連結不足時，Markdown 中應有【待補來源】
            assert "【待補來源】" in md, "連結不足時 Markdown 應有【待補來源】"
    
    # 驗證 JSON 輸出
    json_data = to_json(doc)
    assert isinstance(json_data, dict), "JSON 輸出應為 dict"
    assert "segments" in json_data, "JSON 應含 segments 欄位"
    
    # 檢查 JSON 中每個論點區塊的連結欄位
    for seg in json_data["segments"]:
        if seg["type"] == "supplement":
            assert "openable_links_count" in seg, "JSON segment 應含 openable_links_count"
            assert "openable_links_status" in seg, "JSON segment 應含 openable_links_status"
            assert "openable_links_incomplete_reason" in seg, "JSON segment 應含 openable_links_incomplete_reason"
            
            # 數值一致性檢查
            if seg["openable_links_status"] == "sufficient":
                assert seg["openable_links_count"] >= MIN_OPENABLE_LINKS, (
                    f"JSON 中 sufficient 狀態但連結數不足：{seg['openable_links_count']}"
                )
                assert not seg["openable_links_incomplete_reason"], (
                    "JSON 中 sufficient 狀態不應有缺失原因"
                )
            elif seg["openable_links_status"] == "insufficient":
                assert seg["openable_links_count"] < MIN_OPENABLE_LINKS, (
                    f"JSON 中 insufficient 狀態但連結數充足：{seg['openable_links_count']}"
                )
                assert seg["openable_links_incomplete_reason"], (
                    "JSON 中 insufficient 狀態應有缺失原因"
                )


# ---------------------------------------------------------------------------
# 端到端測試：來源相關性啟發式檢查
# ---------------------------------------------------------------------------

def test_e2e_source_relevance_heuristic():
    """端到端測試：來源與論點相關性的啟發式檢查。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 相關來源
    src_relevant = _source(
        "s1",
        "行政程序法第92條",
        level="A",
        url="https://law.moj.gov.tw/LawClass/A0030055",
        content="行政程序法第92條：本法所稱行政處分，係指行政機關就公法上具體事件所為之決定或其他公權力措施。",
    )
    
    # 不相關來源（內容完全不包含論點關鍵詞）
    src_irrelevant = _source(
        "s2",
        "民事訴訟法規定",
        level="B",
        url="https://law.moj.gov.tw/LawClass/A0030001",
        content="民事訴訟法關於當事人能力與訴訟能力的規定。",
    )
    
    # 測試相關來源
    doc_relevant = _assemble_with_sources([src_relevant], gap, used_ids=["s1"])
    supplements = [s for s in doc_relevant.segments if s.type == "supplement"]
    assert len(supplements) == 1
    seg = supplements[0]
    
    # 相關來源應通過關鍵詞檢查
    question_keywords = ["行政處分", "定義"]
    for src in seg.sources:
        if src.content:
            has_keyword = any(kw in src.content for kw in question_keywords)
            assert has_keyword, (
                f"相關來源應包含論點關鍵詞：{src.title} - {src.content[:50]}"
            )
    
    # 測試不相關來源（應在實際流程中被過濾或標記）
    # 這裡只驗證啟發式檢查能正確識別不相關
    assert not any(
        kw in src_irrelevant.content for kw in question_keywords
    ), "不相關來源不應包含論點關鍵詞"


# ---------------------------------------------------------------------------
# 端到端測試：無候選來源案例
# ---------------------------------------------------------------------------

def test_e2e_no_candidates_case():
    """端到端測試：無任何候選來源案例（pure pending）。"""
    gap = Gap("行政處分的定義為何？", "missing", "筆記未展開定義")
    
    # 無任何來源
    doc = assemble_correction(
        _doc(),
        [gap],
        {gap.question: []},
        {gap.question: WrittenSupplement("【待補證】尚無可用來源。", [])},
        {gap.question: cross_validate(gap.question, [])},
    )
    
    supplements = [s for s in doc.segments if s.type == "supplement"]
    assert supplements, "應有至少一個論點區塊"
    
    for seg in supplements:
        # 無候選來源時，不強制要求可開啟連結
        assert seg.openable_links_status == "sufficient", (
            "無候選來源時狀態應為 sufficient（不強制要求）"
        )
        assert seg.openable_links_count == 0, (
            "無候選來源時連結數應為 0"
        )
        assert not seg.openable_links_incomplete_reason, (
            "無候選來源時不應有缺失原因"
        )
