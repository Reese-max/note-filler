"""為每個被 -m 'not integration' 排除的測試建立 failing-first 對照情境。

任務：證明「未執行該 integration 測試」是否可能掩蓋 correctness 缺陷。
每項對照鎖定被排除測的 determinism 可重現核心；若產品路徑穩定 PASS，
則「因排除而掩蓋的產品缺陷」判 **NOT-REPRODUCIBLE**（見 docs/）。

對照設計原則（failing-first 候選）：
- 現況健康時應 PASS；
- 若對應 correctness 契約被破壞，此測應紅，且預設 CI 因 deselect
  不會跑到 integration 本體——故此對照補上 determinism 可見性。
"""
from __future__ import annotations

import json
import urllib.request
from datetime import date
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from note_filler.domain import detect_domain
from note_filler.gap import Gap, detect_gaps
from note_filler.knowledge.law_lookup import LawLookup
from note_filler.llm import FakeLLM, GrokClient
from note_filler.pipeline import run_pipeline
from note_filler.questions import generate_questions
from note_filler.retrieve import _LEVEL_RANK, retrieve_for_gap
from note_filler.retrieve.models import Source
from note_filler.retrieve.twinkle import TwinkleClient

LAW_DB = Path(__file__).resolve().parents[1] / "data" / "law_index.db"

# 與 tests/deselected_allowlist.json 8 項順序對齊（機器可對照）
EXCLUDED_NODE_IDS: tuple[str, ...] = (
    "tests/test_domain.py::test_detect_domain_real_grok_returns_law",
    "tests/test_e2e_acceptance.py::test_e2e_acceptance_real",
    "tests/test_gap.py::test_detect_gaps_real_grok",
    "tests/test_llm.py::test_grok_pong_integration",
    "tests/test_pipeline.py::test_run_pipeline_real_grok",
    "tests/test_questions.py::test_generate_questions_real_grok",
    "tests/test_retrieve.py::test_retrieve_for_gap_real_twinkle_smoke",
    "tests/test_twinkle.py::test_search_real_twinkle_hub",
)


def test_control_map_covers_all_eight_excluded() -> None:
    """對照模組必須 1:1 覆蓋 8 個 deselected node id（防漏項）。"""
    assert len(EXCLUDED_NODE_IDS) == 8
    assert len(set(EXCLUDED_NODE_IDS)) == 8
    allowlist = json.loads(
        (Path(__file__).resolve().parents[1] / "tests" / "deselected_allowlist.json")
        .read_text(encoding="utf-8")
    )
    assert [row["test_id"] for row in allowlist[:8]] == list(EXCLUDED_NODE_IDS)


# ---------------------------------------------------------------------------
# #1 domain — 排除: test_detect_domain_real_grok_returns_law
# ---------------------------------------------------------------------------
def test_control_01_domain_legal_label_contract() -> None:
    """Failing-first：明顯法律文 + 標籤契約；破壞 strip/fallback 應紅。

    排除測只驗真 Grok 回 law。離線對照鎖定：合法標籤必命中、雜訊不硬猜 law。
    """
    legal = (
        "刑法第271條規定,殺人者處死刑、無期徒刑或十年以上有期徒刑;"
        "前項之未遂犯罰之。本條為普通殺人罪之構成要件與法定刑度。"
    )
    assert detect_domain(legal, FakeLLM(["law"])) == "law"
    assert detect_domain(legal, FakeLLM(["law."])) == "law"
    assert detect_domain(legal, FakeLLM(["  LAW  "])) == "law"
    # 無合法標籤 → other（不得因「看起來像法」就硬猜；硬猜才會掩蓋無來源閘精神）
    assert detect_domain("今天天氣不錯,午餐吃了牛肉麵。", FakeLLM(["嗯..."])) == "other"


# ---------------------------------------------------------------------------
# #2 e2e — 排除: test_e2e_acceptance_real
# ---------------------------------------------------------------------------
@pytest.mark.skipif(not LAW_DB.exists(), reason="缺 data/law_index.db")
def test_control_02_e2e_offline_quality_gates() -> None:
    """Failing-first：離線四硬閘；任一閘回歸應紅（對應 e2e 排除測 determinism 核心）。"""
    fixture = Path(__file__).resolve().parent / "fixtures" / "sample.txt"
    if not fixture.is_file():
        fixture = Path(__file__).resolve().parent / "fixtures" / "real_note.txt"
    note_text = fixture.read_text(encoding="utf-8")
    law = LawLookup(str(LAW_DB))

    class _StubTwinkle:
        def search(self, query: str, n: int = 3) -> list[Source]:
            today = date.today().isoformat()
            return [
                Source(
                    id="stub-a",
                    title="行政程序法 §92",
                    url="https://example.local/law/a",
                    level="A",
                    content="行政處分係行政機關就公法上具體事件所為之決定。",
                    fetched_date=today,
                    doc_date=None,
                    distance=0.1,
                ),
                Source(
                    id="stub-b",
                    title="立法說明",
                    url="https://example.local/law/b",
                    level="B",
                    content="行政程序法第92條修正說明。",
                    fetched_date=today,
                    doc_date=None,
                    distance=0.2,
                ),
            ]

    fake = FakeLLM(
        [
            "law",
            "什麼是行政處分?\n行政程序法第92條的定義為何?\n訴願前置程序為何?",
                json.dumps(
                    [
                        {
                            "question": "什麼是行政處分?",
                            "status": "covered",
                            "reason": "原稿已說明",
                        },
                        {
                            "question": "行政程序法第92條的定義為何?",
                            "status": "missing",
                            "reason": "筆記未展開條文定義",
                        },
                        {
                            "question": "訴願前置程序為何?",
                            "status": "covered",
                            "reason": "原稿已說明",
                        },
                    ],
                ensure_ascii=False,
            ),
            '{"keyword": "行政處分", "law_name": "行政程序法"}',
            "行政處分係指行政機關就公法上具體事件所為之對外發生法律效果之單方行政行為[^1][^2]。",
        ]
    )
    doc = run_pipeline(str(fixture), fake, _StubTwinkle(), law)

    originals = [s.text for s in doc.segments if s.type == "original"]
    assert "".join(originals) == note_text or note_text in "".join(originals) or all(
        part in note_text for part in originals if part.strip()
    )
    # 原稿不可變：original 段必須逐字出現在原稿
    for seg in doc.segments:
        if seg.type == "original":
            assert seg.text in note_text or note_text.find(seg.text) >= 0 or seg.text == note_text
    for seg in doc.segments:
        if seg.type == "supplement" and not seg.sources:
            assert seg.confidence == "pending_evidence"
        if seg.type == "supplement" and seg.confidence == "verified":
            assert seg.sources
            assert "[^" in seg.text


# ---------------------------------------------------------------------------
# #3 gap — 排除: test_detect_gaps_real_grok
# ---------------------------------------------------------------------------
def test_control_03_gap_uncovered_question_must_surface() -> None:
    """Failing-first：明顯未涵蓋題必須進缺口；covered 必須被濾掉。"""
    note = (
        "行政程序法第92條規定,行政處分係行政機關就公法上具體事件所為之"
        "對外直接發生法律效果之單方行政行為。"
    )
    questions = [
        "行政處分的法定定義為何?",
        "行政罰鍰的裁量基準與上限為何?",
    ]
    canned = json.dumps(
        [
            {"question": questions[0], "status": "covered", "reason": "已說明定義"},
            {"question": questions[1], "status": "missing", "reason": "完全沒提罰鍰"},
        ],
        ensure_ascii=False,
    )
    gaps = detect_gaps(questions, note, FakeLLM([canned]))
    assert isinstance(gaps, list)
    assert all(isinstance(g, Gap) for g in gaps)
    assert all(g.status in ("partial", "missing") for g in gaps)
    assert any("裁量基準" in g.question or "罰鍰" in g.question for g in gaps)
    assert not any(g.question == questions[0] and g.status == "covered" for g in gaps)


# ---------------------------------------------------------------------------
# #4 llm — 排除: test_grok_pong_integration
# ---------------------------------------------------------------------------
def test_control_04_grok_client_parse_and_endpoint_contract(monkeypatch: pytest.MonkeyPatch) -> None:
    """Failing-first：GrokClient 必須打對 endpoint 並正確解析 content（PONG 契約的 determinism 層）。"""
    captured: dict = {}

    class FakeResp:
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def read(self):
            return json.dumps(
                {"choices": [{"message": {"content": "PONG"}}]}
            ).encode("utf-8")

    def fake_urlopen(req, timeout=None):
        captured["url"] = req.full_url
        captured["method"] = req.get_method()
        captured["body"] = json.loads(req.data.decode("utf-8"))
        captured["timeout"] = timeout
        return FakeResp()

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    client = GrokClient(api_key="secret", model="grok-4.3", timeout=30)
    out = client.complete(
        [{"role": "user", "content": "Reply with exactly one word: PONG"}],
        temperature=0.0,
    )
    assert "PONG" in out.upper()
    assert captured["url"] == "http://127.0.0.1:8318/v1/chat/completions"
    assert captured["method"] == "POST"
    assert captured["body"]["model"] == "grok-4.3"
    assert captured["timeout"] == 30


# ---------------------------------------------------------------------------
# #5 pipeline — 排除: test_run_pipeline_real_grok
# ---------------------------------------------------------------------------
def test_control_05_pipeline_c6_pending_when_no_sources(tmp_path: Path) -> None:
    """Failing-first：真模型路徑用 empty twinkle 時 C6 仍須守 pending_evidence。"""
    p = tmp_path / "note.docx"
    d = DocxDocument()
    d.add_paragraph("行政程序法要求行政行為應遵守正當程序。")
    d.add_paragraph("本筆記僅記錄部分重點,尚未展開。")
    d.save(str(p))

    class FakeTwinkle:
        def __init__(self) -> None:
            self.queries: list[str] = []

        def search(self, query: str, n: int = 3) -> list[Source]:
            self.queries.append(query)
            return []

    class FakeLaw:
        def lookup_article(self, law_name, article_no):
            return None

        def law_exists(self, name):
            return True

        def fuzzy_find_law(self, name):
            return None

        def search_articles(self, keyword, limit=5, law_name=None):
            return []

    llm = FakeLLM(
        [
            "law",
            "正當程序的要件為何?",
            json.dumps(
                [
                    {
                        "question": "正當程序的要件為何?",
                        "status": "missing",
                        "reason": "未展開",
                    }
                ],
                ensure_ascii=False,
            ),
            '{"keyword": "正當程序", "law_name": null}',
            "【待補證】此問題缺乏可用來源,尚待補充。",
        ]
    )
    twinkle = FakeTwinkle()
    doc = run_pipeline(str(p), llm, twinkle, FakeLaw())
    assert doc.original is not None
    assert isinstance(doc.segments, list)
    for seg in doc.segments:
        if seg.type == "supplement" and not seg.sources:
            assert seg.confidence == "pending_evidence"
            assert seg.text.startswith("【待補證】") or "待補證" in seg.text


# ---------------------------------------------------------------------------
# #6 questions — 排除: test_generate_questions_real_grok
# ---------------------------------------------------------------------------
def test_control_06_questions_clean_list_contract() -> None:
    """Failing-first：問題清單必須非空乾淨字串、不得殘留 JSON 結構符。"""
    note = (
        "個人資料保護法規範公務機關與非公務機關對個人資料之蒐集、處理及利用,"
        "並要求特定情形須告知當事人並取得同意。"
    )
    canned = "個人資料的蒐集要件為何?\n告知義務的例外情形有哪些?"
    result = generate_questions(note, "law", FakeLLM([canned]))
    assert isinstance(result, list)
    assert len(result) >= 1
    assert all(isinstance(q, str) and q.strip() for q in result)
    assert not any(q.strip().startswith(("[", "{")) for q in result)


# ---------------------------------------------------------------------------
# #7 retrieve — 排除: test_retrieve_for_gap_real_twinkle_smoke
# ---------------------------------------------------------------------------
class _EmptyTwinkle:
    def search(self, query: str, n: int = 3) -> list[Source]:
        return []


def _smoke_asserts(out: list[Source]) -> None:
    """逐字重現 deselected #7 smoke（含 vacuous empty 風險）。"""
    assert all(isinstance(s, Source) for s in out)
    assert all(s.level in ("A", "B") for s in out)
    keys = [(_LEVEL_RANK[s.level], s.distance) for s in out]
    assert keys == sorted(keys)


def test_control_07_retrieve_smoke_vacuous_empty_is_blind_spot() -> None:
    """證明 #7 smoke 對 empty 仍綠 → 未執行/弱斷言可掩蓋「檢索全滅」。"""
    empty: list[Source] = []
    _smoke_asserts(empty)
    assert empty == []


@pytest.mark.skipif(not LAW_DB.exists(), reason="缺 data/law_index.db")
def test_control_07b_retrieve_law_level_a_not_vacuous() -> None:
    """Failing-first 產品路徑：law+LawLookup 不得被 vacuous empty 掩蓋。"""
    gap = Gap(question="行政處分附款的容許界限為何?", status="missing", reason="")
    law = LawLookup(str(LAW_DB))
    llm = FakeLLM(['{"keyword": "行政處分", "law_name": "行政程序法"}'])
    out = retrieve_for_gap(gap, "law", _EmptyTwinkle(), law, llm)
    _smoke_asserts(out)
    assert out, (
        "CORRECTNESS：law+LawLookup 回空，但 #7 smoke 對 empty 仍 vacuous PASS；"
        "預設 -m 'not integration' 看不到此缺陷"
    )
    assert any(s.level == "A" for s in out)


# ---------------------------------------------------------------------------
# #8 twinkle — 排除: test_search_real_twinkle_hub
# ---------------------------------------------------------------------------
def test_control_08_twinkle_parsed_source_contract(monkeypatch: pytest.MonkeyPatch) -> None:
    """Failing-first：MCP 解析後 Source 必須 level∈A/B、content 非空、fetched_date 當日。"""
    # 沿用 test_twinkle 既有 SSE 風格最小假回應
    session_headers: list[str] = []

    class FakeResp:
        def __init__(self, body: bytes, headers: dict | None = None):
            self._body = body
            self.headers = headers or {}

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def read(self):
            return self._body

    def fake_urlopen(req, timeout=None):
        method = req.get_method()
        url = req.full_url
        if "initialize" in (req.data or b"").decode("utf-8", errors="ignore") or method == "POST":
            body = req.data.decode("utf-8") if req.data else ""
            if "initialize" in body:
                return FakeResp(
                    b'{"jsonrpc":"2.0","id":1,"result":{"protocolVersion":"2024-11-05"}}',
                    {"mcp-session-id": "sess-1"},
                )
            if "notifications/initialized" in body or '"method": "notifications/initialized"' in body:
                return FakeResp(b"")
            if "tools/call" in body:
                payload = {
                    "jsonrpc": "2.0",
                    "id": 3,
                    "result": {
                        "content": [
                            {
                                "type": "text",
                                "text": json.dumps(
                                    {
                                        "results": [
                                            {
                                                "id": "1120001",
                                                "title": "道路交通管理處罰條例相關提案",
                                                "url": "https://ly.gov.tw/bill/1120001",
                                                "content": "現行條文對累犯之處罰不足，應檢討刑度。",
                                                "date": "2024-01-15",
                                            }
                                        ]
                                    },
                                    ensure_ascii=False,
                                ),
                            }
                        ]
                    },
                }
                return FakeResp(json.dumps(payload).encode("utf-8"), {"mcp-session-id": "sess-1"})
        raise AssertionError(f"unexpected request: {method} {url}")

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    # TwinkleClient 內部 import 路徑可能綁 twinkle 模組的 urllib
    import note_filler.retrieve.twinkle as twinkle_mod

    monkeypatch.setattr(twinkle_mod.urllib.request, "urlopen", fake_urlopen)

    results = TwinkleClient(token="fake-token").search("道路交通管理處罰條例", n=3)
    assert isinstance(results, list)
    assert results, "解析後不得空（對照真 hub 至少能拿到結構）"
    for src in results:
        assert isinstance(src, Source)
        assert src.level in ("A", "B")
        assert src.content.strip(), "全文非空（#8 排除測 determinism 核心）"
        assert src.fetched_date == date.today().isoformat()
    _ = session_headers  # 保留擴充點
