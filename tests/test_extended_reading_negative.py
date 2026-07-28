"""最小負例：延伸閱讀不合格或無關時，流程明確失敗或標示【待補來源】。

兩種案例：
  1. 論點存在但無合格延伸閱讀：補充段產生但無任何候選來源（extended_readings 為空），
     confidence 為 pending_evidence，binding_report 標記 pending_evidence，
     輸出含【待補來源】與缺失綁定原因。
  2. 延伸閱讀可開啟但與論點不直接相關：候選來源有可開啟 URL 但內容與論點無關，
     LLM 以【待補證】表示無法作答，binding_report 仍標記 pending_evidence，
     延伸閱讀列出但不等於合格引用，輸出含【待補來源】。

避免默默產出不完整筆記：每筆缺失綁定皆在 binding_report 與輸出中可見。
"""
from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from note_filler.binding_report import build_binding_report, parse_binding_report
from note_filler.correction import assemble_correction
from note_filler.export import to_docx, to_json, to_markdown
from note_filler.gap import Gap
from note_filler.parse import Document, Paragraph
from note_filler.retrieve.models import Source
from note_filler.verify import cross_validate
from note_filler.write import WrittenSupplement


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _doc(path: str = "input/note.txt") -> Document:
    return Document(
        source_path=path,
        paragraphs=(Paragraph(0, "原稿逐字保留。"),),
        full_text="原稿逐字保留。",
    )


def _source(
    sid: str,
    title: str,
    content: str,
    level: str,
    url: str | None = None,
    distance: float = 0.3,
) -> Source:
    return Source(
        id=sid,
        title=title,
        url=url,
        level=level,
        content=content,
        fetched_date="2026-07-28",
        doc_date=None,
        distance=distance,
    )


def _pending_supplement(text: str, omitted_ids: list[str] | None = None, used_source_ids: list[str] | None = None) -> WrittenSupplement:
    """建立以【待補證】開頭的 WrittenSupplement，可指定 omitted_source_ids 與 used_source_ids。

    pipeline 正常流程會自動補齊 omitted_source_ids；直接呼叫 assemble_correction
    時需手動傳入，模擬 pipeline 行為。
    """
    return WrittenSupplement(
        text=text,
        used_source_ids=used_source_ids or [],
        citation_spans=[],
        omitted_source_ids=omitted_ids or [],
    )


# ---------------------------------------------------------------------------
# 案例一：論點存在但無合格延伸閱讀
# ---------------------------------------------------------------------------

class TestNoQualifiedExtendedReadings:
    """論點存在但無合格延伸閱讀：流程明確標示缺失來源綁定。

    構造：gap 有缺口，retrieved 為空（無候選來源），LLM 回傳【待補證】。
    驗證：
      1. Segment：confidence = pending_evidence, extended_readings = [],
         pending_evidence_reason 非空且指向「檢索無可用來源」
      2. Binding report：binding_status = pending_evidence,
         at_least_one_source = False
      3. Markdown 輸出含「待補證原因」區塊
      4. JSON 輸出的 extended_readings_status 為 pending_evidence
      5. 不得默默產出不含缺失標記的成品（所有輸出格式均含【待補來源】相關標記）
    """

    def test_segment_pending_evidence_with_empty_readings(self):
        """segment 層級：無候選來源時 extended_readings 為空、reason 指出缺口。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        seg = product.segments[-1]
        assert seg.type == "supplement"
        assert seg.confidence == "pending_evidence"
        assert seg.extended_readings == []
        assert seg.extended_readings_status == "pending_evidence"
        assert "檢索無可用來源" in seg.pending_evidence_reason
        assert seg.pending_evidence_reason.strip()

    def test_binding_report_flags_missing_source(self):
        """binding report：無候選來源時 at_least_one_source = False、
        binding_status 為 pending_evidence。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        report = parse_binding_report(build_binding_report(product))
        arg = report["arguments"][0]
        assert arg["checks"]["at_least_one_source"] is False
        # 無候選來源時，intentional_pending 為 True（正確的 pending 行為）
        assert arg["binding_status"] == "pending_evidence"
        assert arg["extended_readings"] == []
        assert arg["extended_readings_status"] == "pending_evidence"
        assert arg["pending_evidence_reason"].strip()
        # binding_ok=True 是正確的：segment 被正確標記為 pending，非結構性失敗

    def test_markdown_output_shows_pending_reason(self):
        """Markdown 輸出：pending_evidence 時含「待補證原因」區塊。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        md = to_markdown(product)
        assert "待補證原因" in md
        assert "檢索無可用來源" in md

    def test_json_output_pending_evidence_status(self):
        """JSON 輸出：extended_readings_status 為 pending_evidence。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        data = to_json(product)
        supplements = [s for s in data["segments"] if s["type"] == "supplement"]
        assert supplements
        seg = supplements[0]
        assert seg["extended_readings_status"] == "pending_evidence"
        assert seg["extended_readings"] == []
        assert seg["pending_evidence_reason"].strip()

    def test_docx_output_includes_pending_marker(self):
        """DOCX 輸出：含缺失來源標記。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        to_docx(product, out_path)
        out_doc = DocxDocument(out_path)
        texts = [p.text for p in out_doc.paragraphs]
        assert any("待補證原因" in t for t in texts)

    def test_original_text_preserved_despite_no_readings(self):
        """無延伸閱讀時原稿逐字不變。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        data = to_json(product)
        originals = [s for s in data["segments"] if s["type"] == "original"]
        for seg in originals:
            assert seg["text"] == "原稿逐字保留。"

    def test_missing_binding_reason_in_report(self):
        """binding report 的 pending_evidence_reason 指出缺失綁定原因。"""
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: []},
            {gap.question: _pending_supplement("【待補證】此問題缺乏可用來源。")},
            {gap.question: cross_validate(gap.question, [])},
        )
        report = parse_binding_report(build_binding_report(product))
        arg = report["arguments"][0]
        assert arg["pending_evidence_reason"].strip()
        reason = arg["pending_evidence_reason"]
        assert "來源" in reason or "候選" in reason or "檢索" in reason


# ---------------------------------------------------------------------------
# 案例二：延伸閱讀可開啟但與論點不直接相關
# ---------------------------------------------------------------------------

class TestExtendedReadingsOpenableButUnrelated:
    """延伸閱讀可開啟但與論點不直接相關：候選來源存在但未被引用。

    構造：retrieved 有候選來源（含有效 URL），但 LLM 因內容與問題「完全無關」
    而以【待補證】作答，used_source_ids 為空 → 候選來源成為延伸閱讀。
    驗證：
      1. Segment：extended_readings_status = available（有候選來源），
         extended_readings 含候選來源，但 confidence = pending_evidence
      2. Binding report：at_least_one_source = False（無實際引用來源），
         binding_status = pending_evidence
      3. Markdown 輸出列出延伸閱讀來源但標示【待補來源】
      4. 延伸閱讀的 URL 可開啟但不等於合格引用
      5. pending_evidence_reason 指出「有候選來源但未被引用」
    """

    def test_extended_readings_available_but_not_cited(self):
        """segment：有候選來源但 LLM 未引用，延伸閱讀含候選來源。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導，與行政程序完全無關。",
            "C",
            url="https://example.com/climate",
            distance=0.9,
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            # omitted_source_ids 模擬 pipeline 行為：未被引用的候選來源
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        seg = product.segments[-1]
        assert seg.type == "supplement"
        assert seg.confidence == "pending_evidence"
        # 候選來源成為延伸閱讀
        assert len(seg.extended_readings) == 1
        assert seg.extended_readings[0]["source_id"] == "web:unrelated"
        assert seg.extended_readings[0]["url"] == "https://example.com/climate"
        assert seg.extended_readings_status == "available"
        assert "有候選來源但未被引用" in seg.pending_evidence_reason
        # 無實際引用來源
        assert seg.sources == []

    def test_binding_report_readings_not_cited(self):
        """binding report：延伸閱讀存在但 at_least_one_source = False。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        report = parse_binding_report(build_binding_report(product))
        arg = report["arguments"][0]
        assert arg["checks"]["at_least_one_source"] is False
        assert arg["binding_status"] == "pending_evidence"
        # 延伸閱讀存在
        assert len(arg["extended_readings"]) == 1
        assert arg["extended_readings"][0]["source_id"] == "web:unrelated"
        assert arg["extended_readings_status"] == "available"

    def test_markdown_lists_unrelated_readings_with_marker(self):
        """Markdown 輸出：列出延伸閱讀但標示【待補來源】。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        md = to_markdown(product)
        # 延伸閱讀區塊存在
        assert "延伸閱讀" in md
        assert "web:unrelated" in md
        # 待補證原因存在
        assert "待補證原因" in md or "待補來源" in md

    def test_json_extended_readings_openable_url(self):
        """JSON 輸出：延伸閱讀含可開啟 URL 但不影響 binding。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        data = to_json(product)
        supplements = [s for s in data["segments"] if s["type"] == "supplement"]
        assert supplements
        seg = supplements[0]
        assert len(seg["extended_readings"]) == 1
        er = seg["extended_readings"][0]
        assert er["source_id"] == "web:unrelated"
        assert er["url"] == "https://example.com/climate"
        # 延伸閱讀有 URL 但不等於合格引用
        assert seg["extended_readings_status"] == "available"
        assert seg["confidence"] == "pending_evidence"

    def test_docx_output_lists_unrelated_readings(self):
        """DOCX 輸出：列出延伸閱讀但標示缺失。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        to_docx(product, out_path)
        out_doc = DocxDocument(out_path)
        texts = [p.text for p in out_doc.paragraphs]
        assert any("延伸閱讀" in t for t in texts)
        assert any("web:unrelated" in t for t in texts)

    def test_markdown_available_status_shows_reason_when_present(self):
        """Markdown 輸出：status=available 但 reason 非空時，待補證原因仍應顯示。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        seg = product.segments[-1]
        assert seg.extended_readings_status == "available"
        assert seg.pending_evidence_reason.strip()
        md = to_markdown(product)
        # available 狀態下延伸閱讀區塊存在
        assert "> **延伸閱讀**" in md
        assert "> - [web:unrelated]" in md
        # 待補證原因仍應顯示（不再只限於 pending_evidence 狀態）
        assert "> **待補證原因**：有候選來源但未被引用" in md

    def test_docx_available_status_shows_reason_when_present(self):
        """DOCX 輸出：status=available 但 reason 非空時，待補證原因仍應顯示。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            out_path = f.name
        to_docx(product, out_path)
        out_doc = DocxDocument(out_path)
        texts = [p.text for p in out_doc.paragraphs]
        # available 狀態下延伸閱讀段落存在
        assert any("延伸閱讀" in t for t in texts)
        assert any("web:unrelated" in t for t in texts)
        # 待補證原因仍應顯示
        assert any("待補證原因" in t for t in texts)
        assert any("有候選來源但未被引用" in t for t in texts)

    def test_pending_evidence_reason_points_to_uncited(self):
        """pending_evidence_reason 指出「有候選來源但未被引用」。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        report = parse_binding_report(build_binding_report(product))
        arg = report["arguments"][0]
        assert arg["pending_evidence_reason"].strip()
        reason = arg["pending_evidence_reason"]
        assert "有候選來源但未被引用" in reason or "來源" in reason

    def test_multiple_unrelated_readings_all_listed(self):
        """多個無關延伸閱讀：全部列出且皆不影響 binding。"""
        src_a = _source(
            "web:climate",
            "氣候變遷報導",
            "全球暖化趨勢。",
            "C",
            url="https://example.com/climate",
        )
        src_b = _source(
            "web:sports",
            "體育新聞",
            "奧運賽事報導。",
            "D",
            url="https://example.com/sports",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a, src_b]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:climate", "web:sports"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        data = to_json(product)
        seg = [s for s in data["segments"] if s["type"] == "supplement"][0]
        assert len(seg["extended_readings"]) == 2
        er_ids = [er["source_id"] for er in seg["extended_readings"]]
        assert "web:climate" in er_ids
        assert "web:sports" in er_ids
        assert seg["extended_readings_status"] == "available"
        assert seg["confidence"] == "pending_evidence"
        # 兩個延伸閱讀 URL 皆可開啟（但不等於合格引用）
        for er in seg["extended_readings"]:
            assert er["url"] is not None

    def test_original_text_preserved_despite_unrelated_readings(self):
        """無關延伸閱讀不影響原稿逐字不變。"""
        unrelated_src = _source(
            "web:unrelated",
            "與行政處分無關的網頁",
            "這是一篇關於氣候變遷的報導。",
            "C",
            url="https://example.com/climate",
        )
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [unrelated_src]},
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["web:unrelated"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        data = to_json(product)
        originals = [s for s in data["segments"] if s["type"] == "original"]
        for seg in originals:
            assert seg["text"] == "原稿逐字保留。"


# ---------------------------------------------------------------------------
# 案例三：延伸閱讀優先級規則驗證
# ---------------------------------------------------------------------------

class TestExtendedReadingsPriorityRules:
    """延伸閱讀優先級規則：實際引用來源優先，再補權威延伸來源（Level A > B > C > D）。

    構造：多個候選來源不同層級，驗證排序與優先級規則。
    驗證：
      1. 延伸閱讀按 Level A > B > C > D 排序
      2. 同層級按 distance 遞增排序
      3. 實際引用來源優先於延伸閱讀
      4. 至少2條URL不足時明確標示【待補來源】
    """

    def test_extended_readings_sorted_by_level_then_distance(self):
        """延伸閱讀按 Level A > B > C > D 排序，同層級按 distance 遞增。"""
        src_a = _source("src_a", "Level A 來源", "A級內容", "A", url="https://a.com", distance=0.8)
        src_b = _source("src_b", "Level B 來源", "B級內容", "B", url="https://b.com", distance=0.3)
        src_c = _source("src_c", "Level C 來源", "C級內容", "C", url="https://c.com", distance=0.5)
        src_d = _source("src_d", "Level D 來源", "D級內容", "D", url="https://d.com", distance=0.2)
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_d, src_b, src_a, src_c]},  # 亂序輸入
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["src_d", "src_b", "src_a", "src_c"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        data = to_json(product)
        seg = [s for s in data["segments"] if s["type"] == "supplement"][0]
        
        # 驗證排序：A > B > C > D
        readings = seg["extended_readings"]
        order = [r["source_id"] for r in readings]
        assert order == ["src_a", "src_b", "src_c", "src_d"], f"排序錯誤: {order}"
        
        # 驗證同層級按 distance 排序（此案例無同層級）

    def test_same_level_sorted_by_distance(self):
        """同層級來源按 distance 遞增排序。"""
        src_b1 = _source("src_b1", "B來源1", "B級內容1", "B", url="https://b1.com", distance=0.7)
        src_b2 = _source("src_b2", "B來源2", "B級內容2", "B", url="https://b2.com", distance=0.3)
        src_b3 = _source("src_b3", "B來源3", "B級內容3", "B", url="https://b3.com", distance=0.5)
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_b3, src_b1, src_b2]},  # 亂序輸入
            {gap.question: _pending_supplement(
                "【待補證】現有來源與問題無關。",
                omitted_ids=["src_b3", "src_b1", "src_b2"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        data = to_json(product)
        seg = [s for s in data["segments"] if s["type"] == "supplement"][0]
        
        # 驗證同層級按 distance 遞增
        readings = seg["extended_readings"]
        order = [r["source_id"] for r in readings]
        assert order == ["src_b2", "src_b3", "src_b1"], f"同層級排序錯誤: {order}"

    def test_cited_sources_priority_over_extended(self):
        """實際引用來源優先於延伸閱讀（在可開啟連結計算中）。"""
        src_a = _source("src_a", "Level A 引用", "A級內容", "A", url="https://a.com", distance=0.4)
        src_b = _source("src_b", "Level B 延伸", "B級內容", "B", url="https://b.com", distance=0.2)
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a, src_b]},
            {gap.question: _pending_supplement(
                "依據 Level A 來源[^1]，行政處分需符合正當程序。",
                omitted_ids=["src_b"],
                used_source_ids=["src_a"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        seg = product.segments[-1]
        
        # 驗證可開啟連結計算：引用來源優先
        assert seg.openable_links_count >= 1, "至少應有引用來源的可開啟連結"
        
        # 驗證延伸閱讀只包含未引用來源
        extended_ids = [r["source_id"] for r in seg.extended_readings]
        assert "src_b" in extended_ids, "延伸閱讀應包含未引用來源"
        assert "src_a" not in extended_ids, "延伸閱讀不應包含已引用來源"
        
        # 驗證實際引用來源存在於 sources
        cited_ids = [s.id for s in seg.sources]
        assert "src_a" in cited_ids, "引用來源應存在於 sources"


# ---------------------------------------------------------------------------
# 案例四：至少2條URL保證驗證
# ---------------------------------------------------------------------------

class TestMinimumTwoURLsRequirement:
    """論點區塊至少需2條真實可開啟URL，不足時明確輸出【待補來源】。

    構造：不同數量的可開啟連結情境。
    驗證：
      1. 可開啟連結 >= 2 條時 status = sufficient
      2. 可開啟連結 = 1 條時 status = insufficient，明確輸出缺口原因
      3. 可開啟連結 = 0 條時 status = insufficient，明確輸出缺口原因
      4. pending_evidence_reason 指出URL數量不足
    """

    def test_sufficient_when_two_or_more_urls(self):
        """可開啟連結 >= 2 條時 status = sufficient。"""
        src_a = _source(
            "src_a", "行政處分法規", "行政處分的法定定義。", "A",
            url="https://a.com", distance=0.3,
        )
        src_b = _source(
            "src_b", "行政處分解釋", "行政處分的要件說明。", "B",
            url="https://b.com", distance=0.4,
        )
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a, src_b]},
            {gap.question: _pending_supplement(
                "依據來源，行政處分需符合正當程序。",
                used_source_ids=["src_a", "src_b"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        seg = product.segments[-1]
        assert seg.openable_links_count >= 2
        assert seg.openable_links_status == "sufficient"
        assert seg.openable_links_incomplete_reason == ""

    def test_insufficient_when_only_one_url(self):
        """可開啟連結 = 1 條時 status = insufficient，明確輸出缺口原因。"""
        src_a = _source("src_a", "A來源", "內容", "A", url="https://a.com", distance=0.3)
        src_b = _source("src_b", "B來源", "內容", "B", url=None, distance=0.4)  # 無URL
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a, src_b]},
            {gap.question: _pending_supplement(
                "依據來源，行政處分需符合正當程序。",
                used_source_ids=["src_a", "src_b"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        seg = product.segments[-1]
        assert seg.openable_links_count == 1
        assert seg.openable_links_status == "insufficient"
        assert "不足 2 條" in seg.openable_links_incomplete_reason
        # 注意：此案例不是 pending_evidence，所以 pending_evidence_reason 不會有 URL 不足訊息

    def test_insufficient_when_no_urls(self):
        """可開啟連結 = 0 條時 status = insufficient，明確輸出缺口原因。"""
        src_a = _source("src_a", "A來源", "內容", "A", url=None, distance=0.3)
        src_b = _source("src_b", "B來源", "內容", "B", url=None, distance=0.4)
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a, src_b]},
            {gap.question: _pending_supplement(
                "依據來源，行政處分需符合正當程序。",
                used_source_ids=["src_a", "src_b"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        seg = product.segments[-1]
        assert seg.openable_links_count == 0
        assert seg.openable_links_status == "insufficient"
        assert "無可開啟連結" in seg.openable_links_incomplete_reason
        # 注意：此案例不是 pending_evidence，所以 pending_evidence_reason 不會有 URL 不足訊息

    def test_pending_evidence_reason_includes_url_count(self):
        """pending_evidence_reason 明確指出URL數量不足。"""
        src_a = _source("src_a", "A來源", "內容", "A", url="https://a.com", distance=0.3)
        
        gap = Gap("行政處分之要件為何？", "missing", "原稿未展開")
        product = assemble_correction(
            _doc(),
            [gap],
            {gap.question: [src_a]},
            {gap.question: _pending_supplement(
                "【待補證】來源不足。",
                used_source_ids=["src_a"],
            )},
            {gap.question: cross_validate(gap.question, [])},
        )
        
        seg = product.segments[-1]
        assert seg.confidence == "pending_evidence"
        # 因為只有1條可開啟連結且是 pending_evidence，所以應該有 URL 不足訊息
        reason = seg.pending_evidence_reason
        has_url_msg = "可開啟連結不足 2 條" in reason
        has_general_msg = "來源不足" in reason
        assert has_url_msg or has_general_msg


# ---------------------------------------------------------------------------
# 案例五：來源綁定品質負例
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    ("unrelated", "expected_count", "reason_fragment"),
    [
        (None, 1, "僅有 1 條可開啟連結"),
        (
            _source(
                "web:climate",
                "全球氣候觀測報告",
                "本報告整理海水溫度、降雨與碳排放趨勢。",
                "C",
                url="https://example.com/climate",
            ),
            2,
            "直接相關僅 1 條",
        ),
    ],
    ids=["有引用但延伸閱讀不足", "可開啟延伸連結與論點無關"],
)
def test_source_binding_quality_negative_cases(unrelated, expected_count, reason_fragment):
    """兩種負例皆降級，且回報 argument_id 下的延伸閱讀綁定缺口。"""
    cited = _source(
        "law:92",
        "行政程序法第92條",
        "行政處分是行政機關就公法具體事件所為的決定。",
        "A",
        url="https://law.moj.gov.tw/LawClass/LawSingle.aspx?pcode=A0030055&flno=92",
    )
    gap = Gap("行政處分的定義為何？", "missing", "原稿未展開行政處分定義")
    sources = [cited] + ([unrelated] if unrelated else [])
    omitted_ids = [unrelated.id] if unrelated else []
    product = assemble_correction(
        _doc(),
        [gap],
        {gap.question: sources},
        {
            gap.question: WrittenSupplement(
                "行政處分是具體公權力決定[^1]。",
                [cited.id],
                omitted_source_ids=omitted_ids,
            )
        },
        {gap.question: cross_validate(gap.question, [cited])},
    )

    seg = product.segments[-1]
    report_arg = parse_binding_report(build_binding_report(product))["arguments"][0]
    reason = report_arg["openable_links_incomplete_reason"]
    markdown = to_markdown(product)

    assert seg.openable_links_count == expected_count
    assert seg.openable_links_status == "insufficient"
    assert [source.id for source in seg.sources] == [cited.id]
    assert [r["source_id"] for r in seg.extended_readings] == omitted_ids
    assert report_arg["checks"]["at_least_two_openable_links"] is False
    assert "argument:0.extended_readings" in reason
    assert reason_fragment in reason
    assert "【待補來源】" in markdown and "argument:0.extended_readings" in markdown
    assert product.original.full_text == "原稿逐字保留。"
