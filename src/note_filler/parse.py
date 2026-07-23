from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Paragraph:
    idx: int
    text: str


@dataclass(frozen=True)
class Document:
    source_path: str
    paragraphs: tuple  # tuple[Paragraph, ...]
    full_text: str


# 空行分段：切在「一或多個只含空白的行」上；段內文字不改寫，僅去除段落
# 首尾空白以確定邊界（原文內容 immutable，不做替換/正規化）。
_BLANK_SEP = re.compile(r"(?:\r?\n)[ \t]*(?:\r?\n)+")


def _split_txt(content: str) -> list[str]:
    blocks = _BLANK_SEP.split(content)
    return [b.strip() for b in blocks if b.strip()]


def _read_docx(path: str) -> list[str]:
    # 延遲 import，讓 .txt 路徑不需安裝 python-docx
    from docx import Document as DocxDocument

    docx = DocxDocument(path)
    return [p.text for p in docx.paragraphs if p.text.strip()]


def parse_note(path: str) -> Document:
    suffix = Path(path).suffix.lower()
    if suffix == ".txt":
        content = Path(path).read_text(encoding="utf-8")
        texts = _split_txt(content)
    elif suffix == ".docx":
        texts = _read_docx(path)
    else:
        raise ValueError(f"不支援的筆記副檔名：{suffix!r}（僅支援 .txt / .docx）")

    paragraphs = tuple(
        Paragraph(idx=i, text=t) for i, t in enumerate(texts)
    )
    full_text = "\n".join(p.text for p in paragraphs)
    if not full_text.strip():
        logger.warning("parse_note: file %s parsed to empty content", path)
    return Document(source_path=path, paragraphs=paragraphs, full_text=full_text)
